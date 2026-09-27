# -*- coding: utf-8 -*-
"""기업 데이터 업로드 양식(엑셀)을 만든다 → junhee/data/templates/AXPORT_기업데이터_업로드양식_v1.xlsx (publish_static.py 가 static/samples/ 로 복사).

시트 구성은 가상 기업 샘플(v0.3)과 같다: 안내 · 기업정보 · 제품정보 · 거래처 · 수출실적 · 물류. 예시 행은 모두 가상이며 '예시' 로 표시한다.
"""
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "templates" / "AXPORT_기업데이터_업로드양식_v1.xlsx"

HEAD_FILL = PatternFill("solid", fgColor="1E3A8A")
HEAD_FONT = Font(name="맑은 고딕", bold=True, color="FFFFFF", size=10)
REQ_FILL = PatternFill("solid", fgColor="FFF7E6")
EX_FONT = Font(name="맑은 고딕", color="6B7280", size=10, italic=True)
BODY_FONT = Font(name="맑은 고딕", size=10)
THIN = Side(style="thin", color="D1D5DB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

SHEETS = {
    "제품정보": {
        "cols": [("제품ID", "필수 · 회사 내부 ID (수출실적·물류에서 참조)", 12), ("모델명", "필수", 16), ("제품군", "선택 · DRAM, NAND, 로직 등", 12), ("입력HSK", "필수 · 10자리 HSK (통제번호 후보 조회)", 14),
                 ("HS코드", "필수 · 6자리 HS (관세 참고치·HS 선택)", 10), ("HS버전", "선택 · HS2022", 10), ("제조국", "선택", 10), ("사양", "선택", 22), ("단위", "선택 · EA", 8), ("단위중량(g)", "선택", 12), ("분류비고", "선택 · 품목분류 확인 필요 사항", 30)],
        "examples": [["P-001", "MD5-16G", "DRAM", "8542321010", "854232", "HS2022", "대한민국", "DDR5 16Gb", "EA", 0.45, ""],
                     ["P-002", "MC-A72", "로직", "8542311000", "854231", "HS2022", "대한민국", "MCU 72MHz", "EA", 0.3, "예시"]],
        "required": ["제품ID", "모델명", "입력HSK", "HS코드"],
    },
    "거래처": {
        "cols": [("거래처ID", "필수 · 회사 내부 ID", 12), ("법인명", "필수 · 영문 법인명 (CSL 제재 명단 검색)", 28), ("별칭", "선택", 16), ("주소", "선택 · CSL 주소 대조", 34), ("도시", "선택", 12), ("국가", "필수 · 한글 국가명", 10), ("국가코드", "필수 · ISO2 (US, DE, VN …)", 10), ("최종사용자관계", "선택 · 최종사용자 / 유통 / 위탁", 14)],
        "examples": [["C-01", "Example Systems Inc.", "예시 시스템즈", "100 Sample St, San Jose, CA", "San Jose", "미국", "US", "최종사용자"],
                     ["C-02", "Muster Elektronik GmbH", "", "Beispielstraße 1, München", "München", "독일", "DE", "유통"]],
        "required": ["거래처ID", "법인명", "국가", "국가코드"],
    },
    "수출실적": {
        "cols": [("실적ID", "필수 · 고유 ID (물류 시트에서 참조)", 12), ("거래일", "필수 · YYYY-MM-DD", 12), ("거래월", "선택 · YYYY-MM (비우면 거래일에서 계산)", 10), ("연도", "선택", 8), ("제품ID", "필수 · 제품정보의 제품ID", 10), ("거래처ID", "필수 · 거래처의 거래처ID", 10),
                 ("목적국", "필수 · 한글 국가명", 10), ("목적국코드", "선택 · ISO2", 10), ("수량", "선택", 10), ("단위", "선택 · EA", 8), ("순중량(kg)", "권장 · kg당 단가 계산", 12), ("금액", "필수 · 숫자 (통화 열의 통화)", 14), ("통화", "필수 · USD", 8), ("인코텀즈", "선택 · FOB, CIF …", 10),
                 ("취소반품", "선택 · Y 면 실제 0 으로 세고 집계 제외", 10), ("적용환율(KRW/USD)", "선택", 16), ("원화환산(원)", "선택 · 참고", 14)],
        "examples": [["R-0001", "2025-01-06", "2025-01", 2025, "P-001", "C-01", "미국", "US", 12000, "EA", 5.4, 52000.0, "USD", "FOB", "N", "", ""],
                     ["R-0002", "2025-01-15", "2025-01", 2025, "P-002", "C-02", "독일", "DE", 8000, "EA", 2.4, 31000.0, "USD", "CIF", "N", "", ""],
                     ["R-0003", "2025-02-03", "2025-02", 2025, "P-001", "C-01", "미국", "US", 0, "EA", 0, 0, "USD", "FOB", "Y", "", "예시: 취소 건"]],
        "required": ["실적ID", "거래일", "제품ID", "거래처ID", "목적국", "금액", "통화"],
    },
    "물류": {
        "cols": [("물류ID", "필수 · 고유 ID", 12), ("실적ID", "필수 · 수출실적의 실적ID", 12), ("운송수단", "필수 · 항공 / 해상", 10), ("출발지", "선택", 12), ("출발지코드", "필수 · 공항 IATA(ICN) 또는 항만 UN/LOCODE(KRPUS)", 12), ("도착지", "선택", 12),
                 ("도착지코드", "필수 · 공항 IATA 또는 항만 UN/LOCODE", 12), ("선적일", "필수 · YYYY-MM-DD", 12), ("예정도착일", "권장 · 납기 준수율", 12), ("실제도착일", "권장 · 납기 준수율·리드타임 (미도착이면 비움)", 12), ("운임(USD)", "권장 · 운송비 비중", 12)],
        "examples": [["L-0001", "R-0001", "항공", "인천공항", "ICN", "로스앤젤레스공항", "LAX", "2025-01-06", "2025-01-09", "2025-01-09", 135.5],
                     ["L-0002", "R-0002", "해상", "부산항", "KRPUS", "함부르크항", "DEHAM", "2025-01-15", "2025-02-19", "2025-02-22", 1240.0]],
        "required": ["물류ID", "실적ID", "운송수단", "출발지코드", "도착지코드", "선적일"],
    },
}
INFO_ROWS = [
    ("template_version", "1.0", "양식 버전 (수정하지 않음)"),
    ("company", "", "필수 · 회사명 (화면에 표시)"),
    ("company_en", "", "선택 · 영문 회사명"),
    ("data_class", "", "필수 · 실제데이터 / 가상데이터"),
    ("기준기간", "", "필수 · 예: 2023-01 ~ 2025-12 (월별)"),
    ("금액 통화", "", "필수 · 예: USD (결제통화) — 참고환율 통화 선택"),
    ("작성일", "", "선택 · YYYY-MM-DD"),
]
GUIDE = [
    ("AXPORT 기업 데이터 업로드 양식 v1", None),
    ("이 파일의 시트(기업정보·제품정보·거래처·수출실적·물류)와 열 이름을 그대로 두고 값만 채워 주세요. 열 순서는 바꿔도 되지만 이름은 바꾸지 않습니다.", None),
    ("", None),
    ("시트", "역할 · 필수 열"),
    ("기업정보", "회사명·결제통화·기준기간 (A열 키, B열 값)"),
    ("제품정보", "제품ID · 모델명 · 입력HSK(10자리) · HS코드(6자리) → 통제번호 후보, 관세 참고치, HS 선택"),
    ("거래처", "거래처ID · 법인명(영문) · 국가 · 국가코드 → 제재 명단(CSL) 검색, 거래처 집중도"),
    ("수출실적", "실적ID · 거래일 · 제품ID · 거래처ID · 목적국 · 금액 · 통화 (순중량은 kg당 단가) → 월별 수출액, 변동, 점수"),
    ("물류", "물류ID · 실적ID · 운송수단 · 출발지코드 · 도착지코드 · 선적일 (+ 예정·실제 도착일, 운임) → 조회 조건, 납기 준수율, 운송비 비중"),
    ("", None),
    ("작성 규칙", None),
    ("1", "날짜는 YYYY-MM-DD, 금액·수량은 숫자만(천 단위 구분 기호 없이), 통화는 USD 를 권장합니다."),
    ("2", "빈칸은 그대로 비워 두세요. 0 이나 임의 값으로 채우지 마세요. 시스템은 빈칸을 '자료 부족'으로 표시하고 해당 계산에서만 제외합니다."),
    ("3", "취소·반품 건은 취소반품 열에 Y 를 적고 수량·금액을 0 으로 둡니다(실제 0). 완전히 같은 행이 두 번 있으면 1개만 남깁니다."),
    ("4", "거래가 없는 달은 행을 넣지 않습니다. 화면에는 '자료 없음'으로 표시되며 0 으로 계산하지 않습니다."),
    ("5", "물류의 실적ID 는 수출실적의 실적ID 와 같아야 연결됩니다. 예정·실제 도착일이 있으면 납기 준수율과 리드타임, 운임(USD)이 있으면 운송비 비중을 계산합니다."),
    ("6", "HS코드가 없는 제품은 HS별 항목(통제번호 후보·관세)에서 제외되고 건수만 표시됩니다."),
    ("7", "업로드 시 결측·오류 목록을 먼저 보여 드립니다. 확인 후 진행하거나 파일을 보완해 다시 올릴 수 있습니다."),
    ("", None),
    ("주의", "파일은 브라우저에서만 확인하며 내용은 서버로 전송하지 않습니다. 시연 환경에서는 등록된 샘플 파일만 분석됩니다. 예시 행(회색 기울임)은 모두 가상 값이므로 지우고 사용하세요."),
]


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row, c)
        cell.font, cell.fill, cell.border = HEAD_FONT, HEAD_FILL, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "안내"
    for i, (a, b) in enumerate(GUIDE, start=1):
        ws.cell(i, 1, a).font = Font(name="맑은 고딕", bold=(b is None and a != ""), size=12 if i == 1 else 10)
        if b is not None:
            ws.cell(i, 2, b).font = BODY_FONT
            ws.cell(i, 2).alignment = Alignment(wrap_text=True, vertical="top")
    style_header(ws, 4, 2)
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 110
    # 기업정보 (키·값)
    wi = wb.create_sheet("기업정보")
    wi.cell(1, 1, "기업 정보 — A열 키는 바꾸지 말고 B열에 값을 채워 주세요").font = Font(name="맑은 고딕", bold=True, size=11)
    wi.cell(3, 1, "키"), wi.cell(3, 2, "값"), wi.cell(3, 3, "설명")
    style_header(wi, 3, 3)
    for i, (k, v, d) in enumerate(INFO_ROWS, start=4):
        wi.cell(i, 1, k).font = BODY_FONT
        wi.cell(i, 2, v).font = BODY_FONT
        wi.cell(i, 2).fill = REQ_FILL if "필수" in d else PatternFill()
        wi.cell(i, 3, d).font = EX_FONT
        for c in range(1, 4):
            wi.cell(i, c).border = BORDER
    wi.column_dimensions["A"].width = 20
    wi.column_dimensions["B"].width = 40
    wi.column_dimensions["C"].width = 52
    # 표 시트
    for name, spec in SHEETS.items():
        w = wb.create_sheet(name)
        cols = spec["cols"]
        for c, (title, desc, width) in enumerate(cols, start=1):
            cell = w.cell(1, c, title)
            cell.comment = Comment(desc, "AXPORT")
            w.column_dimensions[get_column_letter(c)].width = width
            w.cell(2, c, desc).font = EX_FONT
            w.cell(2, c).alignment = Alignment(wrap_text=True, vertical="top")
            w.cell(2, c).border = BORDER
            if title in spec["required"]:
                w.cell(2, c).fill = REQ_FILL
        style_header(w, 1, len(cols))
        w.row_dimensions[2].height = 42
        for r, ex in enumerate(spec["examples"], start=3):
            for c, v in enumerate(ex, start=1):
                cell = w.cell(r, c, v if v != "" else None)
                cell.font = EX_FONT
                cell.border = BORDER
        w.freeze_panes = "A3"
        # 입력값 목록
        if name == "물류":
            dv = DataValidation(type="list", formula1='"항공,해상"', allow_blank=True)
            w.add_data_validation(dv)
            dv.add(f"C3:C2000")
        if name == "수출실적":
            dv1 = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
            dv2 = DataValidation(type="list", formula1='"USD,KRW,EUR,JPY,CNY"', allow_blank=True)
            w.add_data_validation(dv1)
            w.add_data_validation(dv2)
            dv1.add("O3:O5000")
            dv2.add("M3:M5000")
    wb.save(OUT)
    print("written", OUT, OUT.stat().st_size, "bytes")


if __name__ == "__main__":
    main()
