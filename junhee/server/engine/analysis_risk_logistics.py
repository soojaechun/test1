# (junhee) 2026-09-27 sanghyeob/analysis_risk_logistics.py 에서 옮김. 바꾼 것: 패키지 import 경로·원본 자료 경로(junhee/data/raw)만. 계산 로직은 그대로.
"""Evidence-only regulation, logistics and stability analysis.

These readers never execute the legacy scoring pipeline.  A candidate match is
not a legal determination, and observations never become unvalidated scores.
"""
from __future__ import annotations

import csv
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation, localcontext
import hashlib
import io
from pathlib import Path
import re
import struct
import unicodedata
import xml.etree.ElementTree as ET
from zipfile import BadZipFile, ZipFile
import zlib


COUNTRIES = {"US": ("US", "미국", "USA"), "CN": ("CN", "중국", "CHINA"),
             "JP": ("JP", "일본", "JAPAN"), "DE": ("DE", "독일", "GERMANY"),
             "VN": ("VN", "베트남", "VIETNAM")}
VALID_OBSERVATIONS = {"OBSERVED", "OBSERVED_ZERO", "AVAILABLE", "OK"}


def _text(value):
    return "" if value is None else str(value).strip()


def _number(value):
    if value is None or isinstance(value, bool):
        return None
    try:
        result = Decimal(_text(value).replace(",", ""))
        return result if result.is_finite() else None
    except InvalidOperation:
        return None


def _date(value):
    try:
        return date.fromisoformat(_text(value)[:10])
    except ValueError:
        return None


def _hs(value):
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    value = _text(value)
    # Preserve leading zeroes; reject scientific notation and compound codes.
    return value if re.fullmatch(r"[0-9]{6,10}", value) else ""


def _factor(key, label):
    return {"key": key, "label": label, "state": "insufficient", "score": None,
            "score_status": "PENDING_METHODOLOGY", "note": "검증된 원지표와 자료 범위를 표시하며 배점은 보류합니다.",
            "metrics": [], "warnings": [], "checks": [], "sources": [], "series": []}


def _metric(identifier, label, value=None, unit="", period=None, formula="", status=None, reason=None, evidence=None):
    return {"id": identifier, "label": label, "value": value, "unit": unit, "period": period,
            "formula": formula, "status": status or ("OBSERVED" if value is not None else "INSUFFICIENT"),
            "reason": reason, "evidence": evidence or []}


def _evidence(company, row, field):
    return {"source_id": company.get("source", {}).get("id", "company-upload"),
            "sheet": row.get("_sheet"), "row": row.get("_row"), "field": field}


def _add_company_source(factor, company):
    if company.get("source"):
        source = dict(company["source"])
        source.setdefault("provider", "기업 업로드")
        factor["sources"].append(source)


def _read_source(path, provider):
    data = path.read_bytes()
    digest = hashlib.sha256(data).hexdigest()
    return data, {"id": "local-" + digest[:24], "provider": provider, "name": path.name,
                  "sha256": digest, "status": "READ", "retrieved_at": datetime.now(timezone.utc).isoformat(),
                  "temporal_scope": "보유 스냅샷; 조회 시점의 현행성/과거 당시 가용성은 미검증"}


def _csv_rows(data):
    return [{_text(k): v for k, v in row.items() if k is not None}
            for row in csv.DictReader(io.StringIO(data.decode("utf-8-sig")))]


def _scope(company, inputs):
    sheets = company.get("sheets", {})
    products = [row for row in sheets.get("제품정보", [])
                if _hs(row.get("HS코드"))[:6] == inputs.get("hs6")
                and re.sub(r"[^0-9]", "", _text(row.get("HS버전"))) == "2022"]
    product_ids = {_text(row.get("제품ID")) for row in products}
    aliases = COUNTRIES.get(inputs.get("country"), (inputs.get("country"),))
    def selected(row):
        return _text(row.get("제품ID")) in product_ids and _text(row.get("목적국")).upper() in aliases
    planned = [row for row in sheets.get("수출예정거래", []) if selected(row)]
    actual = [row for row in sheets.get("수출실적", []) if selected(row)]
    return products, planned, actual


def _active_actual(rows, as_of):
    result = []
    for row in rows:
        occurred = _date(row.get("거래일"))
        cancelled = _text(row.get("취소반품")).upper() not in ("", "N", "NO", "FALSE", "0", "정상", "아니오")
        if occurred and occurred <= as_of and not cancelled:
            result.append(row)
    return result


def _normal_name(value):
    return " ".join(unicodedata.normalize("NFKC", _text(value)).casefold().split())


def _hwpx_paragraphs(data):
    paragraphs = []
    with ZipFile(io.BytesIO(data)) as package:
        for name in sorted(package.namelist()):
            if re.fullmatch(r"Contents/section[0-9]+\.xml", name):
                root = ET.fromstring(package.read(name))
                for i, paragraph in enumerate(root.findall(".//{*}p"), 1):
                    text = "".join(t.text or "" for t in paragraph.findall("./{*}run/{*}t"))
                    if text.strip():
                        paragraphs.append({"section": name, "paragraph": i, "text": text.strip()})
    return paragraphs


def _hwp_paragraphs(data):
    import olefile
    paragraphs = []
    with olefile.OleFileIO(io.BytesIO(data)) as document:
        header = document.openstream("FileHeader").read()
        flags = struct.unpack_from("<I", header, 36)[0]
        if flags & 2:
            raise ValueError("encrypted HWP")
        for stream in sorted(document.listdir()):
            if len(stream) != 2 or stream[0] != "BodyText" or not stream[1].startswith("Section"):
                continue
            body = document.openstream(stream).read()
            if flags & 1:
                body = zlib.decompress(body, -15)
            offset, index = 0, 0
            while offset + 4 <= len(body):
                record, = struct.unpack_from("<I", body, offset)
                offset += 4
                tag, size = record & 0x3ff, record >> 20
                if size == 0xfff:
                    if offset + 4 > len(body):
                        break
                    size, = struct.unpack_from("<I", body, offset)
                    offset += 4
                if offset + size > len(body):
                    raise ValueError("truncated HWP record")
                if tag == 67:
                    index += 1
                    text = body[offset:offset + size].decode("utf-16le", errors="replace")
                    # Extended inline controls contain seven additional UTF-16 units.
                    text = re.sub(r"[\x01-\x09\x0b\x0c\x0e-\x17].{7}", " ", text, flags=re.S)
                    text = re.sub(r"[\x00-\x1f]", " ", text).strip()
                    if text:
                        paragraphs.append({"section": "/".join(stream), "paragraph": index, "text": text})
                offset += size
    return paragraphs


def evaluate_regulation(company, inputs, root_path, *, law_snapshot=None):
    factor = _factor("regulation", "규제 적합성")
    factor.update(state="review_required", gate_status="REVIEW_REQUIRED", coverage="PARTIAL")
    factor["note"] = "HS·거래처 후보와 보유 원문을 대조한 사전 검토입니다. 수출 허용·금지 또는 허가 불요를 자동 확정하지 않습니다."
    _add_company_source(factor, company)
    base = Path(root_path) / "junhee" / "data" / "raw" / "test1_regulations"
    products, planned, actual = _scope(company, inputs)
    as_of = _date(inputs.get("as_of")) or date.today()
    actual = _active_actual(actual, as_of)
    factor["checks"].append({"label": "기업 제품 연결", "status": "OBSERVED" if products else "MISSING",
                             "detail": f"선택 HS2022 {inputs.get('hs6')}에 연결된 제품 {len(products)}개"})
    hsk_matches, control_codes = [], set()
    try:
        from openpyxl import load_workbook
        data, source = _read_source(base / "HSK연계표_20260901.xlsx", "전략물자 HSK 연계표")
        source["coverage"] = "HSK와 통제번호의 후보 연계; 제품 기술사양에 대한 판정 아님"
        factor["sources"].append(source)
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            worksheet = workbook["HSK연계표"]
            rows = worksheet.iter_rows(values_only=True)
            headers = [_text(value) for value in next(rows)]
            raw_hs = _hs(inputs.get("hs_raw"))
            for row_number, values in enumerate(rows, 2):
                row = dict(zip(headers, values))
                code = _hs(row.get("HSKCD"))
                if not code or not (code == raw_hs if len(raw_hs) == 10 else code[:6] == inputs.get("hs6")):
                    continue
                controls = [_text(c).rstrip(".") for c in _text(row.get("CNTRLNO")).split(",") if _text(c)]
                control_codes.update(controls)
                hsk_matches.append({"hsk": code, "product_name": row.get("HSKNM"), "control_codes": controls,
                                    "match_level": "HSK10" if len(raw_hs) == 10 else "HS6_CANDIDATE",
                                    "evidence": {"source_id": source["id"], "sheet": "HSK연계표", "row": row_number, "field": "CNTRLNO"}})
        finally:
            workbook.close()
        factor["metrics"].append(_metric("hsk_candidates", "HSK 통제번호 연결 후보", len(control_codes), "개",
            formula="선택 HSK10 정확 일치 또는 HS6 하위 HSK의 고유 통제번호 수", status="REVIEW_REQUIRED",
            reason="0개여도 비전략물자 판정이 아닙니다. 사양·용도·최종사용자 검토가 필요합니다.",
            evidence=[m["evidence"] for m in hsk_matches]))
    except (OSError, ValueError, KeyError, ImportError, ET.ParseError, BadZipFile) as exc:
        factor["warnings"].append("HSK 연계표를 읽지 못했습니다: " + type(exc).__name__)
        factor["metrics"].append(_metric("hsk_candidates", "HSK 통제번호 연결 후보", reason="HSK 원본을 읽을 수 없습니다."))
    factor["hsk_candidates"] = hsk_matches
    kotra_matches, unresolved = [], 0
    try:
        path = next(base.glob("대한무역투자진흥공사_*.csv"))
        data, source = _read_source(path, "KOTRA 국별 대세계 수입규제 보유 CSV")
        rows = _csv_rows(data)
        source["record_count"] = len(rows)
        factor["sources"].append(source)
        for record_number, row in enumerate(rows, 1):
            if _text(row.get("규제시행국")).upper() != inputs.get("country"):
                continue
            matched, ambiguous, supplied = [], False, False
            for field, raw in row.items():
                if "HS_코드" not in field or not _text(raw):
                    continue
                supplied = True
                clean = "".join(c for c in unicodedata.normalize("NFKC", _text(raw)) if unicodedata.category(c) != "Cf")
                if _hs(clean):
                    if clean[:6] == inputs.get("hs6"):
                        matched.append({"field": field, "raw": raw})
                else:
                    ambiguous = True
            if ambiguous or not supplied:
                unresolved += 1
            if matched:
                kotra_matches.append({"record_id": row.get("연번"), "destination": row.get("규제시행국"),
                    "product": row.get("품목명"), "measure": row.get("규제형태(진행상황)"),
                    "target_origin": row.get("규제대상국"), "korea_target_flag": row.get("한국대상여부"),
                    "period_text": row.get("최종 판정결과(부과기간)"), "rate_text": row.get("최종 판정결과(관세율)"),
                    "matched_codes": matched, "status": "CANDIDATE_NOT_APPLICABILITY",
                    "evidence": {"source_id": source["id"], "row": record_number, "field": "CSV 데이터 레코드(물리 줄번호 아님)"}})
        factor["metrics"].append(_metric("kotra_candidates", "목적국·HS 수입규제 후보", len(kotra_matches), "건", status="REVIEW_REQUIRED",
            reason="원산지·기업별 예외·현행 효력은 미판정. 한국대상 N인 항목도 원문 후보로만 보존합니다.", evidence=[m["evidence"] for m in kotra_matches]))
        if unresolved:
            factor["warnings"].append(f"선택 목적국 KOTRA 자료 중 {unresolved}개 레코드에 누락·복합/과학표기 HS가 있어 해당 셀의 자동 비교를 보류했습니다.")
    except (OSError, UnicodeError, StopIteration, csv.Error) as exc:
        factor["warnings"].append("KOTRA CSV 조회 불가: " + type(exc).__name__)
        factor["metrics"].append(_metric("kotra_candidates", "목적국·HS 수입규제 후보", reason="KOTRA 원본을 읽을 수 없습니다."))
    factor["regulation_candidates"] = kotra_matches

    party_ids = {_text(row.get("거래처ID")) for row in planned + actual}
    parties = [row for row in company.get("sheets", {}).get("거래처", []) if _text(row.get("거래처ID")) in party_ids]
    names = {}
    for row in parties:
        for field in ("법인명", "별칭"):
            name = _normal_name(row.get(field))
            if name:
                names.setdefault(name, []).append(_evidence(company, row, field))
    for row in planned:
        name = _normal_name(row.get("최종사용자"))
        if name:
            names.setdefault(name, []).append(_evidence(company, row, "최종사용자"))
    csl_matches = []
    try:
        data, source = _read_source(base / "ITA_consolidated_screening_list.csv", "미국 ITA CSL 보유 CSV")
        rows = _csv_rows(data)
        source["record_count"] = len(rows)
        factor["sources"].append(source)
        for record_number, row in enumerate(rows, 1):
            # Only the primary name is automatically compared: alternate-name
            # delimiters are not asserted without a verified schema.
            name = _normal_name(row.get("name"))
            if not name or name not in names:
                continue
            start, end = _date(row.get("start_date")), _date(row.get("end_date"))
            period_status = "EXPIRED_IN_SNAPSHOT" if end and end < as_of else "FUTURE_IN_SNAPSHOT" if start and start > as_of else "EFFECTIVITY_UNVERIFIED"
            csl_matches.append({field: row.get(field) for field in ("_id", "source", "name", "alt_names", "addresses", "ids", "start_date", "end_date", "programs", "license_requirement", "license_policy", "remarks", "source_list_url", "source_information_url")}
                | {"status": "NAME_CANDIDATE", "period_status": period_status,
                   "evidence": [{"source_id": source["id"], "row": record_number, "field": "source + _id + name"}] + names[name]})
        factor["metrics"].append(_metric("csl_candidates", "거래처·최종사용자 이름 일치 후보", len(csl_matches) if names else None, "건",
            status="REVIEW_REQUIRED" if names else "INSUFFICIENT",
            reason="주 이름의 유니코드 정규화 일치만 확인; 별칭·유사명·주소/식별자 실체확인 미완료. 불일치는 제재 해제를 뜻하지 않습니다." if names else "선택 거래의 거래처/최종사용자 이름이 필요합니다.",
            evidence=[item for match in csl_matches for item in match["evidence"]]))
    except (OSError, UnicodeError, csv.Error) as exc:
        factor["warnings"].append("CSL 스냅샷 조회 불가: " + type(exc).__name__)
        factor["metrics"].append(_metric("csl_candidates", "거래처·최종사용자 이름 일치 후보", reason="CSL 원본을 읽을 수 없습니다."))
    factor["party_candidates"] = csl_matches

    excerpts = []
    # Extract readable original paragraphs, explicitly retaining draft status.
    for path in sorted(base.glob("*")):
        if path.suffix.lower() not in (".hwp", ".hwpx"):
            continue
        try:
            data, source = _read_source(path, "보유 전략물자 고시/개정안 원문")
            paragraphs = _hwpx_paragraphs(data) if path.suffix.lower() == ".hwpx" else _hwp_paragraphs(data)
            source.update(legal_status="PROPOSED_OR_UNVERIFIED", coverage="문단 텍스트 일부; 표 배치·주석·조건 연결과 현행 효력 미검증", paragraph_count=len(paragraphs))
            hits = [p for p in paragraphs if any(re.search(r"(?<![A-Za-z0-9])" + re.escape(code) + r"(?![A-Za-z0-9])", p["text"]) for code in control_codes)]
            chosen = hits[:12] if hits else paragraphs[:2]
            source["excerpt"] = [{**p, "text": p["text"][:1800]} for p in chosen]
            factor["sources"].append(source)
            excerpts.append({"source_id": source["id"], "name": path.name, "status": "PROPOSED_OR_UNVERIFIED",
                             "matched_paragraph_count": len(hits), "excerpts": source["excerpt"]})
        except (OSError, ValueError, KeyError, ImportError, ET.ParseError, zlib.error, BadZipFile, struct.error) as exc:
            factor["warnings"].append(path.name + " 원문 추출 불가: " + type(exc).__name__)
    factor["legal_excerpts"] = excerpts
    from .analysis_company_documents import evaluate_company_documents
    documents = evaluate_company_documents(company, inputs)
    factor["document_references"] = documents["references"]
    factor["document_coverage"] = documents["coverage"]
    factor["checks"].extend(documents["checks"])
    factor["warnings"].extend(documents["warnings"])
    factor["metrics"].append(_metric("declared_documents", "대상에 연결된 기업 증빙목록 참조", documents["coverage"]["linked_rows"], "건", status="REVIEW_REQUIRED",
        reason="제품 공통 참고와 거래별 참조를 구분합니다. 목록의 기재 상태만 확인하며 실제 첨부·진위·법적 적용성은 검증하지 않았습니다.",
        evidence=[ref for doc in documents["references"] if doc["linked"] for ref in doc["evidence"]]))
    if law_snapshot is not None:
        factor["sources"].extend(law_snapshot.get("sources", []))
        factor["warnings"].extend(law_snapshot.get("warnings", []))
        factor["checks"].extend(law_snapshot.get("checks", []))
        # Raw source bodies belong only to factor.sources, which public_result
        # strips; do not duplicate the large legal response in this UI summary.
        factor["current_law"] = {key: law_snapshot.get(key) for key in (
            "status", "law", "body", "annexes", "attachments", "lookup_basis", "adapter_version")}
        law = law_snapshot.get("law") or {}
        law_valid = law_snapshot.get("status") in ("OBSERVED", "PARTIAL")
        factor["metrics"].append(_metric("official_law_effective_date", "공식 조회 고시의 시행일", law.get("effective_at") if law_valid else None,
            period=inputs.get("as_of"), status="REFERENCE_ONLY" if law_valid else law_snapshot.get("status", "MISSING"),
            reason="조회 시점의 현행 고시 판본 정보입니다. 해당 제품·거래의 통제 해당 여부와 허가·적합성을 확정하지 않습니다.",
            evidence=law.get("evidence", [])))
    factor["checks"].extend([
        {"label": "한국 전략물자·상황허가", "status": "REVIEW_REQUIRED", "detail": "HSK는 후보 연계입니다. 기술사양과 현행 고시 전체 조건 확인이 필요합니다."},
        {"label": "미국 수출관리·제재", "status": "REVIEW_REQUIRED", "detail": "CSL 이름 후보만 확인. 미국산 함량·ECCN·FDPR·최종사용자/용도 및 거래 구조 검토 필요"},
        {"label": "목적국 수입규제·인증", "status": "REVIEW_REQUIRED", "detail": "KOTRA 보유 CSV는 참고자료입니다. 현행 법령·원산지·규격·인증 적용성 확인 필요"},
        {"label": "원문 효력", "status": "REVIEW_REQUIRED", "detail": "보유 일부개정안·행정예고를 시행 중인 최종 법령으로 취급하지 않습니다."}])
    factor["warnings"].append("미검사 관할·목록·별칭 및 최신 변경이 남아 있어 종합 규제 게이트는 검토 필요입니다.")
    return factor


def _freight_reference(root_path, country, as_of):
    path = (Path(root_path) / "junhee" / "data" / "raw" / "test1_prices" /
            "260915 2026년 8월 해상 수출입 컨테이너 및 항공수입 운송비용 현황.hwpx")
    data, source = _read_source(path, "관세청 운송비용 HWPX 보유 원표")
    paragraphs = _hwpx_paragraphs(data)
    paragraph_text = "\n".join(p["text"] for p in paragraphs)
    if any(marker not in paragraph_text for marker in ("천원/2TEU", "원/kg", "40피트", "항공 수입")):
        raise ValueError("unverified freight units/direction")
    with ZipFile(io.BytesIO(data)) as package:
        root = ET.fromstring(package.read("Contents/section0.xml"))
    tables = root.findall(".//{*}tbl")
    if len(tables) <= 17:
        raise ValueError("freight table layout")
    all_series = []
    expected = [(13, "SEA_EXPORT", 43, 140), (15, "SEA_IMPORT", 43, 140), (17, "AIR_IMPORT", 37, 120)]
    for table_index, direction, row_count, expected_observations in expected:
        rows = tables[table_index].findall("./{*}tr")
        if len(rows) != row_count:
            raise ValueError("freight table row count")
        route, count = "", 0
        for row_index, row in enumerate(rows[1:], 2):
            cells = ["".join(t.text or "" for t in c.findall(".//{*}t")) for c in row.findall("./{*}tc")]
            if "금액" not in cells:
                continue
            index = cells.index("금액")
            if index == 2:
                route = cells[0]
            year = cells[index - 1]
            if year not in ("2025", "2026") or len(cells[index + 1:]) != 12:
                raise ValueError("freight amount layout")
            for month, text in enumerate(cells[index + 1:], 1):
                amount = _number(text)
                if amount is None:
                    if text.strip():
                        raise ValueError("invalid freight amount")
                    continue
                count += 1
                all_series.append({"period": f"{year}-{month:02d}", "route": route, "direction": direction,
                    "value": float(amount), "unit": "천원/2TEU" if direction.startswith("SEA") else "원/kg",
                    "status": "REFERENCE_PROVISIONAL", "evidence": [{"source_id": source["id"], "sheet": "Contents/section0.xml",
                         "row": row_index, "field": f"table[{table_index}], {year}-{month:02d} 금액"}]})
        if count != expected_observations:
            raise ValueError("freight observation count")
    source.update(observed_period="2025-01~2026-08", observed_count=len(all_series),
                  coverage="해상 수출/수입 280건, 항공 수입 120건의 금액 셀 검증. 기업 견적·HS별 통계 아님",
                  published_at="2026-09-15", provisional=True,
                  excerpt=[p for p in paragraphs if any(token in p["text"] for token in ("2TEU =", "할증료", "신고수리일", "수입화물"))])
    routes = {"US": {"미국서부", "미국동부"}, "CN": {"중국"}, "JP": {"일본"}, "VN": {"베트남"}, "DE": {"유럽연합"}}
    # The actual table calls the regional route EU rather than Germany.
    if country == "DE":
        routes[country].update({"EU", "유럽연합(EU)"})
    selected = [p for p in all_series if p["direction"] == "SEA_EXPORT" and p["route"].replace(" ", "") in routes.get(country, set()) and p["period"] < as_of.strftime("%Y-%m")]
    if as_of < date(2026, 9, 15):
        selected = []  # This published snapshot was unavailable at the requested cut-off.
    return source, selected


def evaluate_logistics(company, inputs, root_path):
    factor = _factor("logistics", "물류 여건")
    factor["note"] = "선택 제품·목적국의 기업 일정 관측과 국가별 해상 수출 비용 참고값입니다. 실제 견적·도어투도어 납기·직항 빈도는 별도 검증이 필요합니다."
    _add_company_source(factor, company)
    products, planned, actual = _scope(company, inputs)
    as_of = _date(inputs.get("as_of")) or date.today()
    actual = _active_actual(actual, as_of)
    # The same identifier in both namespaces is ambiguous and must not be guessed.
    transactions = {}
    conflicts = set()
    for row, field in [(r, "거래ID") for r in planned] + [(r, "실적ID") for r in actual]:
        identifier = _text(row.get(field))
        if identifier in transactions:
            conflicts.add(identifier)
        transactions[identifier] = row
    records = [r for r in company.get("sheets", {}).get("물류", []) if _text(r.get("거래ID")) in transactions and _text(r.get("거래ID")) not in conflicts]
    observations, complete, ongoing, invalid, valid_quotes = [], [], 0, 0, []
    for row in records:
        planned_arrival, actual_arrival = _date(row.get("예정도착일")), _date(row.get("실제도착일"))
        actual_departure, planned_departure = _date(row.get("실제출발일")), _date(row.get("예정출발일"))
        if actual_arrival and actual_departure and actual_arrival < actual_departure:
            invalid += 1
            continue
        transaction = transactions[_text(row.get("거래ID"))]
        evidence = [_evidence(company, row, "예정도착일/실제도착일"),
                    _evidence(company, transaction, "제품ID/목적국")]
        done = bool(actual_arrival and actual_arrival <= as_of and planned_arrival and not (actual_departure and actual_departure > as_of))
        if done:
            complete.append((actual_arrival <= planned_arrival, evidence[0]))
        else:
            ongoing += 1
        due = _date(transaction.get("납기일"))
        quote_expiry = _date(row.get("견적유효기간"))
        freight = _number(row.get("운임(USD)"))
        currency = _text(row.get("통화")).upper()
        quote_status = "VALID_DATE_ONLY" if quote_expiry and quote_expiry >= as_of else "EXPIRED" if quote_expiry else "MISSING_EXPIRY"
        if quote_status == "VALID_DATE_ONLY" and freight is not None and freight >= 0 and currency == "USD":
            valid_quotes.append((float(freight), _evidence(company, row, "운임(USD)/통화/견적유효기간")))
        observations.append({"id": row.get("물류ID"), "transaction_id": row.get("거래ID"),
            "origin": row.get("출발지"), "destination": row.get("도착지"), "mode": row.get("운송수단"),
            "planned_departure": row.get("예정출발일"), "planned_arrival": row.get("예정도착일"),
            "actual_arrival": row.get("실제도착일"), "status": "COMPLETED_OBSERVED" if done else "INCOMPLETE_OR_FUTURE",
            "on_time": actual_arrival <= planned_arrival if done else None,
            "planned_segment_calendar_days": (planned_arrival - planned_departure).days if planned_arrival and planned_departure and planned_arrival >= planned_departure else None,
            "delivery_due": transaction.get("납기일"), "planned_arrival_meets_due": planned_arrival <= due if planned_arrival and due else None,
            "freight_usd": float(freight) if freight is not None and freight >= 0 and currency == "USD" else None,
            "quote_status": quote_status, "quote_expiry": row.get("견적유효기간"), "evidence": evidence})
    factor["shipment_observations"] = observations
    factor["metrics"].extend([
        _metric("completed_shipments", "기준일까지 도착 확인된 물류 기록", len(complete), "건", inputs.get("as_of"), evidence=[v[1] for v in complete]),
        _metric("observed_on_time_rate", "완료 기록의 예정 도착 준수율", round(100 * sum(v[0] for v in complete) / len(complete), 4) if complete else None,
                "%", inputs.get("as_of"), "실제도착일≤예정도착일인 완료 기록 / 기준일까지 도착 확인된 기록 ×100",
                reason="일자 단위 표본 관측입니다. 미완료/미기재 기록을 제외하므로 향후 납기 확률이 아닙니다.", evidence=[v[1] for v in complete]),
        _metric("incomplete_shipments", "미완료·도착 미기재 기록", ongoing, "건", inputs.get("as_of")),
        _metric("valid_usd_quote_count", "유효기한이 남은 USD 운임 기재", len(valid_quotes), "건", inputs.get("as_of"),
                reason="서로 다른 화물·노선의 운임을 평균 또는 합산하지 않습니다. 운송사 확정·포함 비용은 미검증.", evidence=[v[1] for v in valid_quotes]),
        _metric("direct_departures_per_week", "직항·무환적 출발 빈도", unit="회/주", reason="출발/도착 허브, 확정 화물 운항표, 환적 여부 및 조회기간 자료가 필요합니다.")])
    from .analysis_company_logistics import evaluate_company_logistics
    delivery = evaluate_company_logistics(company, inputs)
    factor["metrics"].extend(delivery["metrics"])
    factor["delivery_observations"] = delivery["observations"]
    factor["delivery_groups"] = delivery["groups"]
    factor["delivery_coverage"] = delivery["coverage"]
    factor["warnings"].extend(delivery["warnings"])
    factor["checks"].extend(delivery["checks"])
    if invalid or conflicts:
        factor["warnings"].append(f"일정 역전 {invalid}건, 거래ID 네임스페이스 충돌 {len(conflicts)}건을 제외했습니다.")
    try:
        source, reference = _freight_reference(root_path, inputs.get("country"), as_of)
        factor["sources"].append(source)
        factor["series"].extend(reference)
        for route in sorted({r["route"] for r in reference}):
            latest = max((r for r in reference if r["route"] == route), key=lambda r: r["period"])
            factor["metrics"].append(_metric("sea_export_reference_" + route, route + " 해상 수출 비용 참고", latest["value"], latest["unit"], latest["period"],
                status="REFERENCE_PROVISIONAL", reason="HS·기업 개별 화물과 무관한 국가/지역별 총운송비용 잠정 통계. 40피트 컨테이너(2TEU) 기준입니다.", evidence=latest["evidence"]))
        if inputs.get("country") == "DE":
            factor["warnings"].append("해상 운송비는 독일 단독 통계가 아닌 EU 지역 대용 지표입니다.")
        if not reference:
            factor["warnings"].append("기준일에 사용 가능한 선택 국가의 해상 수출 비용 참고값이 없습니다.")
    except (OSError, ValueError, StopIteration, KeyError, ET.ParseError, BadZipFile) as exc:
        factor["warnings"].append("운송비용 HWPX 원표 검증/조회 불가: " + type(exc).__name__)
    factor["warnings"].append("보유 항공 운송비 통계는 수입 방향이므로 항공 수출 견적으로 사용하지 않았습니다.")
    factor["checks"].append({"label": "기업 물류 연결", "status": "OBSERVED" if records else "MISSING",
                             "detail": f"HS2022 {inputs.get('hs6')} / {inputs.get('country')}의 물류 {len(records)}건 연결"})
    factor["state"] = "partial" if records or factor["series"] else "insufficient"
    return factor


def _month_offset(period, offset):
    year, month = map(int, period.split("-"))
    index = year * 12 + month - 1 + offset
    return f"{index // 12:04d}-{index % 12 + 1:02d}"


def _stability_window(trade, completed_month):
    """Validate the server-selected common window without choosing a new one.

    A missing key keeps the original calendar-window behavior for old callers.
    A present but unavailable/invalid window is diagnostic-only: it must not
    quietly become another independently chosen comparison period.
    """
    if "stability_window" not in trade:
        return completed_month, True, None
    supplied = trade["stability_window"]
    window = dict(supplied) if isinstance(supplied, dict) else {}
    end = window.get("end_month")
    valid_month = isinstance(end, str) and re.fullmatch(r"[0-9]{4}-(?:0[1-9]|1[0-2])", end)
    valid_policy = (window.get("status") in ("ready", "insufficient")
                    and isinstance(window.get("eligible"), bool)
                    and type(window.get("max_lag_months")) is int
                    and window["max_lag_months"] == 3)
    if valid_month:
        year, month = map(int, end.split("-"))
        completed_year, completed = map(int, completed_month.split("-"))
        lag = (completed_year - year) * 12 + completed - month
        valid_lag = (0 <= lag <= 3 and type(window.get("lag_months")) is int
                     and window["lag_months"] == lag)
        if valid_policy and valid_lag:
            return end, True, window
    reason = window.get("reason") if end is None and valid_policy else None
    window.update(end_month=None, eligible=False, status="insufficient", lag_months=None,
                  max_lag_months=3,
                  reason=reason or "공통 안정성 기간이 없거나 월/지연 정책 검증을 통과하지 못했습니다.")
    return completed_month, False, window


def evaluate_stability(company, inputs, trade, root_path):
    factor = _factor("stability", "안정성")
    factor["note"] = "선택 목적국·HS의 월간 대세계 수입액 36개월을 사용합니다. 기업 매출 변동과 시장 변동을 구분하며 점수는 보류합니다."
    _add_company_source(factor, company)
    factor["sources"].extend(trade.get("sources", []))
    as_of = _date(inputs.get("as_of")) or date.today()
    completed_month = _month_offset(as_of.strftime("%Y-%m"), -1)
    end, calculation_allowed, comparison_window = _stability_window(trade, completed_month)
    if comparison_window is not None:
        factor["comparison_window"] = comparison_window
        if calculation_allowed:
            factor["warnings"].append(
                f"공통 안정성 평가 종료월 {end}, 최근 완료월 대비 {comparison_window['lag_months']}개월 지연. "
                + ("선택국의 비교 자료요건 충족." if comparison_window["eligible"] else "선택국은 이 공통기간의 비교 자료요건 미충족."))
        else:
            factor["warnings"].append("공통 평가기간 미확정으로 CV·급감 지표를 보류했습니다. 관측 범위/시계열은 최근 완료월 기준의 진단용입니다.")
        if comparison_window.get("reason"):
            factor["warnings"].append(_text(comparison_window["reason"]))
    periods = [_month_offset(end, offset) for offset in range(-35, 1)]
    indexed, duplicates = {}, set()
    for row in trade.get("monthly_imports", []):
        period = row.get("period")
        if period in indexed:
            duplicates.add(period)
        indexed[period] = row
    values, evidence = [], []
    for period in periods:
        row = indexed.get(period, {})
        amount = _number(row.get("value_decimal", row.get("value")))
        if amount is not None and row.get("value_decimal") is not None:
            display_amount = _number(row.get("value"))
            if display_amount is None or float(amount) != float(display_amount):
                amount = None
        status = _text(row.get("status")).upper()
        valid = (period not in duplicates and amount is not None and amount >= 0
                 and status in VALID_OBSERVATIONS and (status != "OBSERVED_ZERO" or amount == 0))
        if not valid and status in VALID_OBSERVATIONS:
            status = "INVALID_DATA"
        value = amount if valid else None
        values.append(value)
        factor["series"].append({"period": period, "value": float(value) if value is not None else None,
            "status": ("OBSERVED_ZERO" if value == 0 else "OBSERVED") if valid else "DUPLICATE_PERIOD" if period in duplicates else status or "MISSING",
            "unit": "USD", "evidence": row.get("evidence", []), "reason": row.get("reason")})
        if valid:
            evidence.extend(row.get("evidence", []))
    n = sum(value is not None for value in values)
    full = n == 36
    average = sum(values) / 36 if full else None
    cv = None
    if calculation_allowed and full and average > 0:
        with localcontext() as context:
            context.prec = 50
            mean = sum(values) / Decimal(36)
            variance = sum((value - mean) ** 2 for value in values) / Decimal(35)
            cv = float(variance.sqrt() / mean)
    changes, drops = [], 0
    for index in range(1, 36):
        previous, current = values[index - 1], values[index]
        if previous is None or current is None or previous <= 0:
            continue
        dropped = current <= previous * Decimal("0.8")
        drops += dropped
        changes.append({"period": periods[index], "previous_period": periods[index - 1],
                        "change_pct": float((current / previous - 1) * 100), "drop_at_least_20pct": dropped})
    factor["monthly_changes"] = changes
    coverage = f"{periods[0]}~{periods[-1]}"
    complete_pairs = len(changes) == 35
    comparison_eligible = bool(comparison_window and calculation_allowed
                               and comparison_window["status"] == "ready" and comparison_window["eligible"]
                               and full and complete_pairs and cv is not None)
    if comparison_window is not None:
        factor["comparison_eligible"] = comparison_eligible
        factor["checks"].append({"label": "공통기간 비교 가능 여부",
            "status": "OBSERVED" if comparison_eligible else "NOT_COMPARABLE",
            "detail": (f"기간 {coverage} · 공통창 {comparison_window.get('window_id', '미확정')} · "
                       f"기간정책 {comparison_window.get('policy_version', '미확인')} · "
                       + ("비교 자료요건 충족; 원지표와 상대 참고지수의 근거를 별도로 표시합니다." if comparison_eligible
                          else "공통 비교 자료요건 또는 최소 2개국 요건 미충족. 관측값은 개별 참고이며 점수는 보류합니다."))})
        if comparison_window.get("eligible") and not (full and complete_pairs and cv is not None):
            factor["warnings"].append("공통창의 비교 가능 표시와 선택국 원관측 검사가 일치하지 않아 비교 가능으로 표시하지 않았습니다.")
    metric_status = "OBSERVED" if full and complete_pairs else "PARTIAL"
    if not calculation_allowed:
        metric_status = "WINDOW_UNAVAILABLE"
    factor["metrics"].extend([
        _metric("market_months", "월간 수입 관측 범위", n, "개월/36개월", coverage,
                reason="동일 HS2022·목적국·USD의 완료된 달만 포함. 누락 월을 0으로 채우지 않습니다.", evidence=evidence),
        _metric("market_cv_36m", "36개월 월간 수입 변동계수", cv, "비율", coverage,
                "표본표준편차(ddof=1) / 36개월 평균 수입액",
                status="WINDOW_UNAVAILABLE" if not calculation_allowed else None,
                reason="공통 평가기간 미확정. 표시 기간은 시계열 진단용이며 36개월 CV를 산출하지 않습니다." if not calculation_allowed else "36개 연속 관측값과 양수 평균이 필요합니다." if cv is None else "계절성·추세를 제거하지 않은 원시 월간 수입액의 변동성입니다.", evidence=evidence),
        _metric("market_drop_pairs", "급감 계산에 유효한 인접 월 쌍", len(changes), "쌍/35쌍", coverage,
                "이전 달 수입액>0이고 두 인접 달이 모두 관측된 경우", evidence=evidence),
        _metric("market_drop_count", "전월 대비 20% 이상 감소", drops if calculation_allowed and changes else None, "회", coverage,
                "현재 달≤이전 달×0.8; 두 관측값이 인접하며 이전 달>0",
                status=metric_status, reason="공통 평가기간이 없어 급감 횟수를 보류합니다." if not calculation_allowed else "유효한 월 쌍만 센 관측 횟수입니다. 누락·이전 달 0은 비급감으로 처리하지 않습니다.", evidence=evidence),
        _metric("market_drop_frequency", "급감 빈도(유효 쌍 기준)", drops / len(changes) if calculation_allowed and changes else None, "비율", coverage,
                "20% 이상 감소 횟수 / 유효한 인접 월 쌍 수", status=metric_status,
                reason="공통 평가기간이 없어 급감 빈도를 보류합니다." if not calculation_allowed else "36개월/35쌍이 완비되어야 전체 기간 지표로 해석할 수 있습니다.", evidence=evidence)])
    if not full:
        factor["warnings"].append(f"표시 기간 {coverage}의 36개월 중 {36-n}개월이 누락/오류/중복되어 36개월 CV를 보류했습니다.")
    if full and average == 0:
        factor["warnings"].append("36개월 수입액이 모두 0입니다. 무역 관측 없음이며 안정성 우수 판정이 아닙니다.")
    if values[-1] == 0:
        factor["warnings"].append("최근 기준월의 수입액이 실제 0으로 관측되었습니다. 상대 변동성 점수가 시장 수요의 존재를 보장하지 않습니다.")
    if not complete_pairs:
        factor["warnings"].append(f"급감 빈도 분모는 35쌍 중 {len(changes)}쌍입니다. 0 기준값 또는 결측을 건너뛰어 연결하지 않습니다.")
    _, planned, actual = _scope(company, inputs)
    actual = _active_actual(actual, as_of)
    currencies = sorted({_text(r.get("통화")).upper() for r in planned + actual if _text(r.get("통화"))})
    factor["settlement_currencies"] = currencies
    factor["metrics"].append(_metric("settlement_currency_count", "선택 거래의 기재 통화", len(currencies), "종류", inputs.get("as_of"),
        reason=", ".join(currencies) if currencies else "통화가 기재된 선택 거래가 없습니다.", evidence=[_evidence(company, r, "통화") for r in planned + actual]))
    fx = trade.get("fx")
    if fx and fx.get("metrics"):
        factor["metrics"].extend(m for m in fx["metrics"] if not m.get("id", "").startswith("fx_reference_rate:"))
        factor["sources"].extend(fx.get("sources", []))
        factor["checks"].extend(fx.get("checks", []))
        factor["warnings"].extend(fx.get("warnings", []))
        factor["fx_series"] = fx.get("series", [])
        factor["calendar_status"] = fx.get("calendar_status", "UNVERIFIED")
    else:
        factor["metrics"].append(_metric("fx_volatility_60d", "결제통화 환율 60거래일 변동성", unit="비율",
            formula="61개 환율 수준으로 산출한 60개 일간 로그수익률의 표본표준편차",
            reason="결제통화별 통일 기준통화·61개 검증 관측·거래일 달력이 필요합니다. 보유 FRED 월변동/휴일 미검증 일별 자료를 이 지표로 대체하지 않습니다."))
        factor["warnings"].extend((fx or {}).get("warnings", []))
    diagnosis = trade.get("missing_diagnostics", {})
    factor["checks"].extend({"label": f'{d["period"]} 수입 자료 확인',
                             "status": d["diagnostic_status"], "detail": d["reason"]}
                            for d in diagnosis.get("diagnostics", []) if d["period"] in periods)
    factor["checks"].append({"label": "시장 안정성 관측 완결성", "status": "OBSERVED" if full and complete_pairs and cv is not None else "PARTIAL",
                             "detail": f"{coverage}, 유효 월 {n}/36, 유효 인접 쌍 {len(changes)}/35"})
    factor["state"] = "observed" if full and complete_pairs and cv is not None else "partial" if n else "insufficient"
    return factor
