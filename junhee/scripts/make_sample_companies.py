# -*- coding: utf-8 -*-
"""업로드 시연용 가상 기업 2개의 수출데이터 엑셀을 junhee/data/samples/ 에 만든다.

- 실존 기업명·실제 재무수치를 쓰지 않는다. 모든 수치는 시연용으로 임의 작성한 가상값이다.
- 시트 구성은 AXPORT_초기개발명세.md §5 표준 시트 제안을 그대로 따른다.
- 매크로·수식 없음 (값만 기록).
- 검증 로직 테스트용 오류 케이스를 일부러 섞고, '검증케이스' 시트에 위치를 적는다.

실행: python junhee/scripts/make_sample_companies.py
"""
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "samples"
OUT.mkdir(parents=True, exist_ok=True)

TEMPLATE_VERSION = "0.1"
NOTICE = "시연용 가상 데이터. 실제 기업 자료가 아님"

SHEETS = {
    "제품정보": ["제품ID", "모델명", "제품군", "HS코드", "HS버전", "제조국", "사양", "사양서참조"],
    "수출예정거래": ["거래ID", "제품ID", "목적국", "거래처ID", "최종사용자", "용도", "수량", "단위", "단가(USD)", "통화", "납기일", "거래조건"],
    "수출실적": ["실적ID", "거래일", "제품ID", "거래처ID", "목적국", "수량", "단위", "금액", "통화", "취소반품"],
    "원가·비용": ["제품ID", "제품원가(단위당)", "통화", "운송비(단위당)", "보험료(단위당)", "포장비(단위당)", "부담주체", "결제조건"],
    "재고·생산": ["제품ID", "가용재고", "예약재고", "기존수주량", "생산일정(월)", "생산예정수량", "검사기간(일)", "공급가능수량"],
    "물류": ["물류ID", "거래ID", "출발지", "도착지", "운송수단", "예정출발일", "실제출발일", "예정도착일", "운임(USD)", "통화", "견적유효기간"],
    "거래처": ["거래처ID", "법인명", "별칭", "주소", "국가", "등록번호", "최종사용자관계"],
    "증빙목록": ["증빙ID", "관련제품ID", "관련거래ID", "문서종류", "문서번호", "발행일", "만료일", "첨부참조"],
}
CASE_HEADER = ["케이스", "시트", "엑셀행번호", "식별자", "설명"]


def d(y, m, dd):
    return date(y, m, dd)


# ---------------------------------------------------------------- 회사 1: 한빛반도체 (규모 큼, DRAM 중심 + NAND)
HANBIT = {
    "file": "한빛반도체_수출데이터_v0.1.xlsx",
    "company": "한빛반도체",
    "제품정보": [
        ["P-HB-001", "HBD5-16G", "DRAM", "854232", "HS2022", "대한민국", "DDR5 16Gb, 6400Mbps", "SPEC-HB-001.pdf"],
        ["P-HB-002", "HBD5-24G", "DRAM", "854232", "HS2022", "대한민국", "DDR5 24Gb, 7200Mbps", "SPEC-HB-002.pdf"],
        ["P-HB-003", "HBL-HBM3", "DRAM(HBM)", "854232", "HS2022", "대한민국", "HBM3 12단 24GB", "SPEC-HB-003.pdf"],
        ["P-HB-004", "HBN-1T", "NAND", "854232", "HS2022", "대한민국", "TLC 1Tb 238단", "SPEC-HB-004.pdf"],
        ["P-HB-005", "HBC-SSD2T", "SSD(컨트롤러 IC)", "854231", "HS2022", "대한민국", "PCIe 5.0 2TB", "SPEC-HB-005.pdf"],
        ["P-HB-006", "HBP-MCU", "프로세서", "854231", "HS2022", "대한민국", "32bit MCU 저전력", "SPEC-HB-006.pdf"],
    ],
    "수출예정거래": [
        ["T-HB-001", "P-HB-001", "미국", "C-HB-01", "노스리버 시스템즈", "서버 메모리 모듈", 400000, "EA", 6.20, "USD", d(2026, 11, 15), "FOB"],
        ["T-HB-002", "P-HB-003", "미국", "C-HB-01", "노스리버 시스템즈", "AI 가속기 탑재", 60000, "EA", 210.00, "USD", d(2026, 12, 10), "CIF"],
        ["T-HB-003", "P-HB-004", "중국", "C-HB-02", "화룬 스토리지", "소비자용 SSD", 900000, "EA", 4.10, "USD", d(2026, 10, 30), "FOB"],
        ["T-HB-004", "P-HB-002", "영국", "C-HB-03", "브리스톨 컴퓨트", "데이터센터", 120000, "EA", 9.30, "USD", d(2027, 1, 20), "CIF"],
        ["T-HB-005", "P-HB-005", "인도", "C-HB-04", "델리 디지털", "노트북 조립", 50000, "EA", None, "USD", d(2026, 12, 5), "FOB"],            # 빈칸: 단가
        ["T-HB-006", "P-HB-006", "멕시코", "C-HB-99", "미등록 거래처", "가전 제어", 200000, "EA", 1.85, "USD", d(2027, 2, 1), "FOB"],      # 연결 오류: 거래처ID 없음
        ["T-HB-007", "P-HB-001", "캐나다", "C-HB-05", "메이플 서버", "서버 메모리", 80000, "EA", 6.35, "USD", d(2026, 11, 28), "DAP"],
    ],
    # 수출실적: 2025-09 ~ 2026-09 (13개월) 을 고정 시드로 생성 + 오류 케이스 행 (아래 ACTUALS 참조)
    "actuals_spec": {
        "prefix": "HB",
        "seed": 20260923,
        "products": [  # (제품ID, 월 기준 수량, 기준 단가 USD, [(거래처ID, 국가), ...])
            ("P-HB-001", 300000, 6.20, [("C-HB-01", "미국"), ("C-HB-05", "캐나다")]),
            ("P-HB-002", 90000, 9.30, [("C-HB-03", "영국")]),
            ("P-HB-003", 30000, 210.00, [("C-HB-01", "미국")]),
            ("P-HB-004", 700000, 4.10, [("C-HB-02", "중국")]),
            ("P-HB-005", 25000, 105.00, [("C-HB-04", "인도")]),
            ("P-HB-006", 120000, 1.85, [("C-HB-06", "멕시코")]),
        ],
        "per_month": (2, 3),
        "growth_per_month": 0.03,
        "zero_row": [d(2026, 8, 14), "P-HB-001", "C-HB-05", "캐나다", 0, "EA", 0, "USD", "Y"],
        "bad_date_row": ["2026-13-05", "P-HB-004", "C-HB-02", "중국", 200000, "EA", 820000.00, "USD", "N"],
        "dup_row": [d(2026, 8, 28), "P-HB-006", "C-HB-06", "멕시코", 150000, "EA", 277500.00, "USD", "N"],
    },
    "원가·비용": [
        ["P-HB-001", 4.10, "USD", 0.08, 0.02, 0.03, "수출자", "T/T 30일"],
        ["P-HB-002", 6.40, "USD", 0.08, 0.02, 0.03, "수출자", "T/T 30일"],
        ["P-HB-003", 145.00, "USD", 1.20, 0.60, 0.40, "수출자", "L/C at sight"],
        ["P-HB-004", 2.90, "USD", 0.05, 0.01, 0.02, "수입자", "T/T 45일"],
        ["P-HB-005", 68.00, "USD", 0.90, 0.30, 0.50, "수출자", "T/T 60일"],
        ["P-HB-006", 1.10, "USD", 0.03, 0.01, 0.01, "수입자", "T/T 30일"],
    ],
    "재고·생산": [
        ["P-HB-001", 520000, 400000, 350000, "2026-10", 900000, 7, 1070000],
        ["P-HB-002", 150000, 120000, 100000, "2026-11", 300000, 7, 330000],
        ["P-HB-003", 45000, 60000, 40000, "2026-11", 90000, 14, 75000],
        ["P-HB-004", 1200000, 900000, 800000, "2026-10", 1500000, 5, 1800000],
        ["P-HB-005", 40000, 50000, 30000, "2026-12", 120000, 10, 110000],
        ["P-HB-006", 300000, 200000, 150000, "2026-10", 400000, 5, 500000],
    ],
    "물류": [
        ["L-HB-001", "T-HB-001", "인천공항", "샌프란시스코", "항공", d(2026, 11, 10), None, d(2026, 11, 12), 38000, "USD", d(2026, 10, 31)],
        ["L-HB-002", "T-HB-002", "인천공항", "샌프란시스코", "항공", d(2026, 12, 5), None, d(2026, 12, 7), 42000, "USD", d(2026, 11, 20)],
        ["L-HB-003", "T-HB-003", "부산항", "상하이항", "해상", d(2026, 10, 22), None, d(2026, 10, 26), 5200, "USD", d(2026, 10, 15)],
        ["L-HB-004", "T-HB-004", "부산항", "사우샘프턴항", "해상", d(2027, 1, 2), None, d(2027, 2, 6), 6100, "USD", d(2026, 12, 20)],
        ["L-HB-005", "T-HB-007", "부산항", "밴쿠버항", "해상", d(2026, 11, 12), None, d(2026, 11, 30), 7800, "USD", d(2026, 11, 1)],
    ],
    "거래처": [
        ["C-HB-01", "North River Systems Inc.", "노스리버 시스템즈", "1200 Harbor Blvd, San Jose, CA", "미국", "US-EIN-77-0000001", "최종사용자"],
        ["C-HB-02", "Huarun Storage Co., Ltd.", "화룬 스토리지", "88 Keyuan Rd, Shenzhen", "중국", "CN-91440300000000X", "최종사용자"],
        ["C-HB-03", "Bristol Compute Ltd.", "브리스톨 컴퓨트", "5 Temple Way, Bristol", "영국", "GB-00000001", "유통업체"],
        ["C-HB-04", "Delhi Digital Pvt. Ltd.", "델리 디지털", "Sector 62, Noida", "인도", "IN-U00000DL0000PTC000001", "최종사용자"],
        ["C-HB-05", "Maple Server Corp.", "메이플 서버", "200 Bay St, Toronto", "캐나다", "CA-000000001", "최종사용자"],
        ["C-HB-06", "Monterrey Electro S.A.", "몬테레이 일렉트로", "Av. Constitución 500, Monterrey", "멕시코", "MX-MEL000001AAA", "유통업체"],
    ],
    "증빙목록": [
        ["D-HB-001", "P-HB-001", "T-HB-001", "원산지증명서", "CO-2026-0101", d(2026, 9, 1), d(2027, 8, 31), "docs/CO-2026-0101.pdf"],
        ["D-HB-002", "P-HB-003", "T-HB-002", "전략물자 판정서", "JG-2026-0042", d(2026, 8, 20), d(2028, 8, 19), "docs/JG-2026-0042.pdf"],
        ["D-HB-003", "P-HB-004", "T-HB-003", "상업송장", "INV-2026-0330", d(2026, 9, 10), None, "docs/INV-2026-0330.pdf"],
        ["D-HB-004", "P-HB-002", "T-HB-004", "수출허가서", "EL-2026-0007", d(2026, 9, 5), d(2027, 9, 4), "docs/EL-2026-0007.pdf"],
        ["D-HB-005", "P-HB-005", "T-HB-005", "사양서", "SPEC-HB-005", d(2026, 3, 1), None, "docs/SPEC-HB-005.pdf"],
    ],
}

# ---------------------------------------------------------------- 회사 2: 대성일렉트로닉스 (규모 작음, NAND 중심)
DAESUNG = {
    "file": "대성일렉트로닉스_수출데이터_v0.1.xlsx",
    "company": "대성일렉트로닉스",
    "제품정보": [
        ["P-DS-001", "DSN-512", "NAND", "854232", "HS2022", "대한민국", "TLC 512Gb 176단", "SPEC-DS-001.pdf"],
        ["P-DS-002", "DSN-1T", "NAND", "854232", "HS2022", "대한민국", "QLC 1Tb 176단", "SPEC-DS-002.pdf"],
        ["P-DS-003", "DSD4-8G", "DRAM", "854232", "HS2022", "대한민국", "DDR4 8Gb, 3200Mbps", "SPEC-DS-003.pdf"],
        ["P-DS-004", "DSC-eMMC", "eMMC(컨트롤러 IC)", "854231", "HS2022", "대한민국", "eMMC 5.1 64GB", "SPEC-DS-004.pdf"],
    ],
    "수출예정거래": [
        ["T-DS-001", "P-DS-001", "중국", "C-DS-01", "광저우 메모리텍", "USB 저장장치", 300000, "EA", 1.95, "USD", d(2026, 10, 25), "FOB"],
        ["T-DS-002", "P-DS-002", "미국", "C-DS-02", "레이크사이드 스토리지", "외장 SSD", 120000, "EA", 3.60, "USD", d(2026, 11, 20), "CIF"],
        ["T-DS-003", "P-DS-004", "인도", "C-DS-03", "뭄바이 모바일", "스마트폰 조립", 250000, "EA", 2.40, "USD", d(2026, 12, 1), "FOB"],
        ["T-DS-004", "P-DS-003", "멕시코", "C-DS-04", "티후아나 어셈블리", "가전 제어 보드", 60000, "EA", None, "USD", d(2026, 12, 15), "FOB"],   # 빈칸: 단가
        ["T-DS-005", "P-DS-001", "스위스", "C-DS-77", "미등록 거래처", "산업용 장비", 20000, "EA", 2.05, "USD", d(2027, 1, 10), "DAP"],     # 연결 오류: 거래처ID 없음
    ],
    "actuals_spec": {
        "prefix": "DS",
        "seed": 20260924,
        "products": [
            ("P-DS-001", 200000, 1.95, [("C-DS-01", "중국")]),
            ("P-DS-002", 80000, 3.60, [("C-DS-02", "미국")]),
            ("P-DS-003", 40000, 2.10, [("C-DS-04", "멕시코")]),
            ("P-DS-004", 150000, 2.40, [("C-DS-03", "인도")]),
        ],
        "per_month": (1, 2),
        "growth_per_month": 0.015,
        "zero_row": [d(2026, 8, 9), "P-DS-003", "C-DS-04", "멕시코", 0, "EA", 0, "USD", "Y"],
        "bad_date_row": ["08/21/2026", "P-DS-001", "C-DS-01", "중국", 100000, "EA", 195000.00, "USD", "N"],
        "dup_row": [d(2026, 9, 2), "P-DS-002", "C-DS-02", "미국", 50000, "EA", 180000.00, "USD", "N"],
    },
    "원가·비용": [
        ["P-DS-001", 1.40, "USD", 0.03, 0.01, 0.01, "수출자", "T/T 30일"],
        ["P-DS-002", 2.60, "USD", 0.04, 0.01, 0.02, "수출자", "T/T 45일"],
        ["P-DS-003", 2.10, "USD", 0.03, 0.01, 0.01, "수입자", "T/T 30일"],
        ["P-DS-004", 1.70, "USD", 0.03, 0.01, 0.01, "수출자", "L/C at sight"],
    ],
    "재고·생산": [
        ["P-DS-001", 350000, 300000, 250000, "2026-10", 500000, 5, 600000],
        ["P-DS-002", 90000, 120000, 100000, "2026-11", 200000, 5, 170000],
        ["P-DS-003", 70000, 60000, 50000, "2026-11", 100000, 7, 110000],
        ["P-DS-004", 200000, 250000, 200000, "2026-10", 300000, 5, 250000],
    ],
    "물류": [
        ["L-DS-001", "T-DS-001", "부산항", "광저우항", "해상", d(2026, 10, 18), None, d(2026, 10, 23), 1800, "USD", d(2026, 10, 10)],
        ["L-DS-002", "T-DS-002", "부산항", "롱비치항", "해상", d(2026, 11, 8), None, d(2026, 11, 24), 6900, "USD", d(2026, 10, 30)],
        ["L-DS-003", "T-DS-003", "인천공항", "뭄바이공항", "항공", d(2026, 11, 26), None, d(2026, 11, 27), 9500, "USD", d(2026, 11, 15)],
    ],
    "거래처": [
        ["C-DS-01", "Guangzhou MemoryTech Co., Ltd.", "광저우 메모리텍", "Tianhe District, Guangzhou", "중국", "CN-91440100000000Y", "최종사용자"],
        ["C-DS-02", "Lakeside Storage LLC", "레이크사이드 스토리지", "400 Lake Ave, Chicago, IL", "미국", "US-EIN-36-0000002", "유통업체"],
        ["C-DS-03", "Mumbai Mobile Pvt. Ltd.", "뭄바이 모바일", "Andheri East, Mumbai", "인도", "IN-U00000MH0000PTC000002", "최종사용자"],
        ["C-DS-04", "Tijuana Assembly S.A.", "티후아나 어셈블리", "Blvd. Industrial 900, Tijuana", "멕시코", "MX-TAS000002BBB", "최종사용자"],
    ],
    "증빙목록": [
        ["D-DS-001", "P-DS-001", "T-DS-001", "원산지증명서", "CO-2026-0201", d(2026, 9, 3), d(2027, 9, 2), "docs/CO-2026-0201.pdf"],
        ["D-DS-002", "P-DS-002", "T-DS-002", "상업송장", "INV-2026-0410", d(2026, 9, 12), None, "docs/INV-2026-0410.pdf"],
        ["D-DS-003", "P-DS-004", "T-DS-003", "사양서", "SPEC-DS-004", d(2026, 2, 10), None, "docs/SPEC-DS-004.pdf"],
    ],
}

ACTUALS_START = (2025, 9)   # 시작 연월
ACTUALS_MONTHS = 13         # 2025-09 ~ 2026-09


def build_actuals(spec):
    """고정 시드로 월별 수출실적을 만든다. 반환: (행 목록, 오류 케이스 ID dict)"""
    import random

    rng = random.Random(spec["seed"])
    prefix = spec["prefix"]
    rows, n = [], 0
    y0, m0 = ACTUALS_START
    lo, hi = spec["per_month"]
    for k in range(ACTUALS_MONTHS):
        yy, mm = y0 + (m0 - 1 + k) // 12, (m0 - 1 + k) % 12 + 1
        picks = rng.sample(spec["products"], k=min(len(spec["products"]), rng.randint(lo, hi)))
        for pid, base_qty, base_price, customers in sorted(picks):
            n += 1
            cust, country = rng.choice(customers)
            qty = int(round(base_qty * rng.uniform(0.7, 1.3) * (1 + spec["growth_per_month"] * k), -3))
            price = round(base_price * rng.uniform(0.93, 1.07), 2)
            rows.append([f"R-{prefix}-{n:03d}", d(yy, mm, rng.randint(3, 27)), pid, cust, country,
                         qty, "EA", round(qty * price, 2), "USD", "N"])
    rows.sort(key=lambda r: r[1])
    ids = {}
    n += 1; ids["zero"] = f"R-{prefix}-{n:03d}"; rows.append([ids["zero"]] + spec["zero_row"])
    n += 1; ids["bad_date"] = f"R-{prefix}-{n:03d}"; rows.append([ids["bad_date"]] + spec["bad_date_row"])
    n += 1; ids["dup"] = f"R-{prefix}-{n:03d}"; rows.append([ids["dup"]] + spec["dup_row"]); rows.append([ids["dup"]] + spec["dup_row"])
    return rows, ids


def build_cases(company, ids):
    fixed = {
        "한빛반도체": [
            ("빈칸", "수출예정거래", "T-HB-005", "단가(USD) 셀이 비어 있음 (임의 값·0 으로 채우면 안 됨)"),
            ("연결 오류", "수출예정거래", "T-HB-006", "거래처ID C-HB-99 가 거래처 시트에 없음 (제품ID P-HB-006 은 존재)"),
        ],
        "대성일렉트로닉스": [
            ("빈칸", "수출예정거래", "T-DS-004", "단가(USD) 셀이 비어 있음 (임의 값·0 으로 채우면 안 됨)"),
            ("연결 오류", "수출예정거래", "T-DS-005", "거래처ID C-DS-77 가 거래처 시트에 없음 (제품ID P-DS-001 은 존재)"),
        ],
    }[company]
    bad = "거래일 '2026-13-05' — 존재하지 않는 월" if company == "한빛반도체" else "거래일 '08/21/2026' — 표준(YYYY-MM-DD)이 아닌 문자열"
    return [
        ("값 0", "수출실적", ids["zero"], "취소/반품 건: 수량 0, 금액 0 (0 과 빈칸을 구분해야 함)"),
        *fixed,
        ("날짜 형식 오류", "수출실적", ids["bad_date"], bad),
        ("중복 행", "수출실적", ids["dup"], "동일 내용 행이 2회 기록됨"),
    ]

BOLD = Font(bold=True)
NOTE_FILL = PatternFill("solid", fgColor="FFF4CC")


def write_sheet(wb, name, header, rows, first=False, company=None):
    ws = wb.create_sheet(name)
    r = 1
    if first:
        ws.cell(r, 1, NOTICE).font = BOLD
        ws.cell(r, 1).fill = NOTE_FILL
        r += 1
        ws.cell(r, 1, "template_version")
        ws.cell(r, 2, TEMPLATE_VERSION)
        r += 1
        ws.cell(r, 1, "company")
        ws.cell(r, 2, company)
        r += 1
        ws.cell(r, 1, "data_class")
        ws.cell(r, 2, "가상데이터")
        r += 2
    header_row = r
    for c, h in enumerate(header, 1):
        ws.cell(r, c, h).font = BOLD
    for row in rows:
        r += 1
        for c, v in enumerate(row, 1):
            cell = ws.cell(r, c, v)
            if isinstance(v, date):
                cell.number_format = "YYYY-MM-DD"
    for c in range(1, len(header) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 18
    ws.freeze_panes = ws.cell(header_row + 1, 1)
    return header_row


def build(spec):
    spec = dict(spec)
    spec["수출실적"], case_ids = build_actuals(spec["actuals_spec"])
    cases = build_cases(spec["company"], case_ids)
    wb = Workbook()
    wb.remove(wb.active)
    header_rows = {}
    for i, (name, header) in enumerate(SHEETS.items()):
        header_rows[name] = write_sheet(wb, name, header, spec[name], first=(i == 0), company=spec["company"])

    # 검증케이스 시트: 식별자 → 실제 엑셀 행번호를 찾아 기록
    ws = wb.create_sheet("검증케이스")
    ws.cell(1, 1, f"{spec['company']} — 검증 로직 테스트용 오류 케이스 위치").font = BOLD
    for c, h in enumerate(CASE_HEADER, 1):
        ws.cell(3, c, h).font = BOLD
    r = 3
    for case, sheet, ident, desc in cases:
        hr = header_rows[sheet]
        ids = [row[0] for row in spec[sheet]]
        rownums = [hr + 1 + i for i, v in enumerate(ids) if v == ident]
        r += 1
        ws.cell(r, 1, case)
        ws.cell(r, 2, sheet)
        ws.cell(r, 3, ", ".join(str(n) for n in rownums))
        ws.cell(r, 4, ident)
        ws.cell(r, 5, desc)
    for c, w in zip("ABCDE", (16, 16, 12, 12, 70)):
        ws.column_dimensions[c].width = w

    out = OUT / spec["file"]
    wb.save(out)
    return out


def main():
    for spec in (HANBIT, DAESUNG):
        out = build(spec)
        print("saved", out.relative_to(ROOT.parent), f"({out.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
