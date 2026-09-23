# -*- coding: utf-8 -*-
"""junhee/data/raw 의 원본에서 대시보드 5개 항목 요약값을 추출해 JSON 으로 저장한다.

- 원본 파일은 읽기만 한다 (수정 금지).
- 항목별 결과: junhee/data/processed/<key>.json
- 통합 결과  : junhee/data/processed/dashboard_summary.json
- 자료가 없거나 파싱에 실패한 항목은 state="insufficient" 로 두고 사유를 남긴다.
  값을 지어내지 않는다.
- HS 는 HS_LIST 에 있는 코드를 모두 본다 (854231 프로세서·컨트롤러, 854232 메모리).

실행: python junhee/scripts/extract_summary.py
"""
import calendar
import csv
import hashlib
import json
import re
import sys
import warnings
import zipfile
from datetime import date, datetime
from pathlib import Path

import openpyxl

warnings.filterwarnings("ignore")
ROOT = Path(__file__).resolve().parents[1]  # junhee/
RAW = ROOT / "data" / "raw"
OUT = ROOT / "data" / "processed"
OUT.mkdir(parents=True, exist_ok=True)

TARGET_ISO2 = "US"
TARGET_NAME = "미국"
HS_LIST = ["854231", "854232"]
HS_NAMES = {"854231": "프로세서·컨트롤러", "854232": "메모리"}
HS_LABEL = "·".join(HS_LIST)

F_CSL = RAW / "test1_regulations" / "ITA_consolidated_screening_list.csv"
F_KOTRA = RAW / "test1_regulations" / "대한무역투자진흥공사_국별 대세계 수입규제 현황_20260603.csv"
F_HSK = RAW / "test1_regulations" / "HSK연계표_20260901.xlsx"
F_WSTS = RAW / "test1_marketability" / "WSTS-Historical-Billings-Report-Jul_2026.xlsx"
F_WSTS_STAB = RAW / "test1_stablity" / "WSTS-Historical-Billings-Report-Jul_2026.xlsx"
F_PRICE_IDX = RAW / "test1_prices" / "(통계)2026년08월수출입물가지수무역지수(잠정).xlsx"
F_TARIFF_US = RAW / "test1_prices" / "C840_C410.csv"
F_FREIGHT = RAW / "test1_prices" / "260915 2026년 8월 해상 수출입 컨테이너 및 항공수입 운송비용 현황.hwpx"


def md5(path):
    h = hashlib.md5()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def file_date(path):
    return datetime.fromtimestamp(path.stat().st_mtime).date().isoformat()


def month_end(y, m):
    return date(y, m, calendar.monthrange(y, m)[1]).isoformat()


def ymd_from_name(path, ext):
    d = re.search(r"_(\d{8})\." + ext + "$", path.name).group(1)
    return f"{d[:4]}-{d[4:6]}-{d[6:]}"


def insufficient(key, source, missing, reason, **extra):
    """자료 부족 항목. headline 에 숫자를 넣지 않는다."""
    item = {
        "key": key,
        "headline": "자료 부족",
        "note": missing,
        "state": "insufficient",
        "source": source,
        "as_of": None,
        "data_class": "실제자료",
        "reason": reason,
        "checked_on": date.today().isoformat(),
    }
    item.update(extra)
    return item


# ---------------------------------------------------------------- 규제
def extract_regulation():
    # (1) KOTRA 수입규제: 미국 시행 규제 중 한국 대상 건수, 각 HS 포함 여부
    with open(F_KOTRA, encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))
    hdr = list(rows[0].keys())
    hs_cols = [h for h in hdr if "HS_코드" in h]

    def hs_hit(r, prefix):
        return any(r[c].strip().startswith(prefix) for c in hs_cols if r[c].strip())

    us_rows = [r for r in rows if r["규제시행국"].strip() == TARGET_ISO2]
    us_kr = [r for r in us_rows if r["한국대상여부"].strip() == "Y"]
    per_hs = {
        hs: {
            "rows_us_targeting_korea": sum(1 for r in us_kr if hs_hit(r, hs)),
            "rows_any_country": sum(1 for r in rows if hs_hit(r, hs)),
        }
        for hs in HS_LIST
    }
    in_list = {hs: per_hs[hs]["rows_us_targeting_korea"] > 0 for hs in HS_LIST}
    kotra_as_of = ymd_from_name(F_KOTRA, "csv")

    # (2) CSL 총 건수 · 기준일 (기준일 컬럼이 없어 최신 등재일 + 파일 확보일을 함께 기록)
    with open(F_CSL, encoding="utf-8-sig", newline="") as fh:
        csl = list(csv.DictReader(fh))
    iso_dates = [r["start_date"] for r in csl if re.fullmatch(r"\d{4}-\d{2}-\d{2}", r["start_date"] or "")]
    csl_latest = max(iso_dates) if iso_dates else None

    # (3) HSK 연계표: 각 HS 하위 HSK 코드의 전략물자 통제번호 연계 (사실 확인만, 판정 아님)
    wb = openpyxl.load_workbook(F_HSK, read_only=True, data_only=True)
    hsk_all = [r for r in wb.worksheets[0].iter_rows(values_only=True) if r and r[0]]
    hsk_links = {
        hs: [{"hsk": str(r[0]), "name": r[1], "control_numbers": str(r[3])} for r in hsk_all if str(r[0]).startswith(hs)]
        for hs in HS_LIST
    }
    hsk_as_of = ymd_from_name(F_HSK, "xlsx")

    hit_hs = [hs for hs in HS_LIST if in_list[hs]]
    headline = f"HS {HS_LABEL} 수입규제 " + (f"{'·'.join(hit_hs)} 포함" if hit_hs else "해당 없음")
    note = (
        f"{TARGET_NAME}의 대한국 수입규제 {len(us_kr)}건(KOTRA) · CSL {len(csl):,}건(최신 등재 {csl_latest})"
        " · 전략물자 별표1~4 원문 확보, 판정 미적용"
    )
    annex_files = sorted(p.name for p in (RAW / "test1_regulations").iterdir() if p.name.startswith("[별표"))
    return {
        "key": "regulation",
        "headline": headline,
        "note": note,
        "state": "ok",
        "source": F_KOTRA.name,
        "as_of": kotra_as_of,
        "data_class": "실제자료",
        "detail": {
            "target_country": TARGET_NAME,
            "hs": HS_LIST,
            "kotra": {
                "file": F_KOTRA.name,
                "as_of": kotra_as_of,
                "rows_total": len(rows),
                "rows_us": len(us_rows),
                "rows_us_targeting_korea": len(us_kr),
                "per_hs": per_hs,
                "hs_in_import_regulation_list": in_list,
            },
            "csl": {
                "file": F_CSL.name,
                "rows_total": len(csl),
                "latest_start_date": csl_latest,
                "file_obtained": file_date(F_CSL),
                "note": "파일에 기준일 컬럼이 없어 최신 등재일과 파일 확보일을 기록",
            },
            "hsk_linkage": {
                "file": F_HSK.name,
                "as_of": hsk_as_of,
                "hsk_codes_per_hs": {hs: len(v) for hs, v in hsk_links.items()},
                "links": hsk_links,
                "note": "HSK 코드가 전략물자 통제번호와 연계돼 있다는 사실만 표시. 해당 여부는 별표 원문 대조 후 판정",
            },
            "strategic_items_annex": {
                "status": "규제 원문 확보, 판정 미적용",
                "files": annex_files,
                "note": "별표1~4(HWP·PDF, 614쪽)는 원문 대조 미완료로 이번에 파싱하지 않음",
            },
        },
    }


# ---------------------------------------------------------------- 시장성
def read_wsts_worldwide(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Monthly Data"]
    series = {}  # (year, month) -> value (1000 US$)
    year = None
    for r in ws.iter_rows(values_only=True):
        if r and isinstance(r[0], int) and 1980 <= r[0] <= 2100:
            year = r[0]
        elif r and r[0] == "Worldwide" and year:
            for m in range(12):
                v = r[1 + m]
                if isinstance(v, (int, float)):
                    series[(year, m + 1)] = float(v)
    return series


def extract_market():
    series = read_wsts_worldwide(F_WSTS)
    (y, m) = max(series)
    latest = series[(y, m)]
    prev_year = series.get((y - 1, m))
    prev_month = series.get((y, m - 1) if m > 1 else (y - 1, 12))
    yoy = (latest / prev_year - 1) * 100 if prev_year else None
    mom = (latest / prev_month - 1) * 100 if prev_month else None
    latest_busd = latest / 1e6  # 1000 US$ -> billion US$
    latest_eok_usd = latest / 1e5  # 1000 US$ -> 억 US$
    period = f"{y}-{m:02d}"
    headline = f"{latest_eok_usd:,.0f}억 달러 · 전년동월비 {yoy:+.1f}%"
    note = f"세계 반도체 월간 출하액(WSTS Worldwide) {period} · 전월비 {mom:+.1f}% · 품목(HS) 구분 없는 전체 반도체"
    return {
        "key": "market",
        "headline": headline,
        "note": note,
        "state": "ok",
        "source": F_WSTS.name,
        "as_of": month_end(y, m),
        "data_class": "실제자료",
        "detail": {
            "region": "Worldwide",
            "period": period,
            "sheet": "Monthly Data",
            "unit_raw": "1000 US$",
            "latest_value_thousand_usd": latest,
            "latest_value_billion_usd": round(latest_busd, 3),
            "prev_year_same_month_thousand_usd": prev_year,
            "prev_month_thousand_usd": prev_month,
            "yoy_pct": round(yoy, 2) if yoy is not None else None,
            "mom_pct": round(mom, 2) if mom is not None else None,
            "hs_breakdown_available": False,
            # 종합 탭 스파크라인용 최근 13개월 (실제 값)
            "monthly": [
                {"month": f"{yy}-{mm:02d}", "value_thousand_usd": series[(yy, mm)]}
                for (yy, mm) in sorted(series)[-13:]
            ],
        },
    }


# ---------------------------------------------------------------- 가격
def extract_price():
    # (1) 미국(C840)의 대한국(C410) HS 별 적용관세율·수입액 — 최신 조치일 기준
    with open(F_TARIFF_US, encoding="utf-8-sig", newline="") as fh:
        rows = [r for r in csv.DictReader(fh) if r["hs_code"] in HS_LIST]
    by_hs = {}
    for hs in HS_LIST:
        hs_rows = sorted((r for r in rows if r["hs_code"] == hs), key=lambda r: r["year_dt"])
        if not hs_rows:
            raise RuntimeError(f"HS {hs} row not found in {F_TARIFF_US.name}")
        latest = hs_rows[-1]
        by_hs[hs] = {
            "name": HS_NAMES.get(hs, ""),
            "action_date_latest": latest["year_dt"],
            "best_avlbl_pct": float(latest["best_avlbl"]),
            "imports_usd": float(latest["imports"]),
            "tariff_history": [
                {"date": d, "best_avlbl_pct": v} for d, v in sorted({(r["year_dt"], float(r["best_avlbl"])) for r in hs_rows})
            ],
        }
    any_row = rows[0]
    tariffs = sorted({v["best_avlbl_pct"] for v in by_hs.values()})
    imports_total = sum(v["imports_usd"] for v in by_hs.values())
    latest_date = max(v["action_date_latest"] for v in by_hs.values())
    tariff_txt = f"{tariffs[0]:.1f}%" if len(tariffs) == 1 else f"{tariffs[0]:.1f}~{tariffs[-1]:.1f}%"

    # (2) 수출물가지수: 파일에 '반도체' 세부 항목이 있으면 그 행, 없으면 상위 분류를 표시하고 그 사실을 기록
    wb = openpyxl.load_workbook(F_PRICE_IDX, read_only=True, data_only=True)
    ws = wb["1.수출(기본분류)"]
    header_periods = None
    parent_row = None
    semi_row = None
    for r in ws.iter_rows(values_only=True):
        if r and isinstance(r[2], str) and re.match(r"\d{4}\. ?\d{1,2}", r[2]):
            header_periods = (r[2].strip(), str(r[3]).strip())
        if r and isinstance(r[0], str):
            label = r[0].replace(" ", "")
            if "반도체" in label and semi_row is None:
                semi_row = r
            if label == "컴퓨터,전자및광학기기":
                parent_row = r
    if semi_row is not None:
        idx_row, idx_note = semi_row, "반도체 세부 지수"
    elif parent_row is not None:
        idx_row, idx_note = parent_row, "파일에 '반도체' 세부 지수가 없어 상위 분류(컴퓨터,전자및광학기기)를 표시"
    else:
        raise RuntimeError("수출물가지수 행을 찾지 못함")
    idx_label = idx_row[0]
    prev_v, cur_v, mom, ytd, yoy = idx_row[2], idx_row[3], idx_row[4], idx_row[5], idx_row[6]
    m = re.search(r"(\d{4})년(\d{2})월", F_PRICE_IDX.name)
    idx_period = f"{m.group(1)}-{m.group(2)}"

    headline = f"적용관세 {tariff_txt} · 수입액 합 {imports_total / 1e8:,.1f}억 달러"
    per_hs_txt = " / ".join(f"{hs} {v['imports_usd'] / 1e8:,.1f}억" for hs, v in by_hs.items())
    lead = f"{TARGET_NAME}의 대한국 수입 {per_hs_txt} 달러(조치일 {latest_date})"
    note = (
        f"{lead} · 수출물가(컴퓨터·전자·광학) {cur_v}, {idx_period}p 전월비 {mom:+.1f}% · 반도체 세부지수 미제공"
        if semi_row is None
        else f"{lead} · 반도체 수출물가 {cur_v}, {idx_period}p 전월비 {mom:+.1f}%"
    )
    return {
        "key": "price",
        "headline": headline,
        "note": note,
        "state": "ok",
        "source": F_TARIFF_US.name,
        "as_of": latest_date,
        "data_class": "실제자료",
        "detail": {
            "tariff": {
                "file": F_TARIFF_US.name,
                "reporter": any_row["reporter_name"],
                "partner": any_row["partner_name"],
                "hs": HS_LIST,
                "per_hs": by_hs,
                "imports_total_usd": imports_total,
                "note": "best_avlbl = 조치 시행 후 적용 예상 관세율(%), imports = 최신 가용 수입액(USD)",
            },
            "export_price_index": {
                "file": F_PRICE_IDX.name,
                "sheet": "1.수출(기본분류)",
                "row_label": idx_label,
                "base": "2020=100, 원화기준",
                "periods": header_periods,
                "prev": prev_v,
                "latest_preliminary": cur_v,
                "mom_pct": mom,
                "ytd_pct": ytd,
                "yoy_pct": yoy,
                "period": idx_period,
                "semiconductor_row_available": semi_row is not None,
                "note": idx_note,
            },
        },
    }


# ---------------------------------------------------------------- 물류
def hwpx_text_lines(path):
    with zipfile.ZipFile(path) as z:
        xml = z.read("Contents/section0.xml").decode("utf-8")
    paras = re.findall(r"<hp:p[ >].*?</hp:p>", xml, flags=re.S)
    lines = []
    for p in paras:
        t = "".join(re.findall(r"<hp:t[^>]*>(.*?)</hp:t>", p, flags=re.S))
        t = re.sub(r"<[^>]+>", "", t).strip()
        if t:
            lines.append(t)
    return lines


def extract_logistics():
    src = F_FREIGHT.name
    try:
        lines = hwpx_text_lines(F_FREIGHT)
    except Exception as e:  # noqa: BLE001
        return insufficient(
            "logistics", src,
            "자료 부족 — hwpx 본문을 읽지 못함. 운임 표를 CSV/XLSX 로 별도 확보 필요",
            f"hwpx parse error: {e!r}",
        )

    press = next((l for l in lines if re.match(r"\d{4}\. \d{1,2}\. \d{1,2}\.", l)), None)
    press_date = None
    if press:
        y, mo, d = re.match(r"(\d{4})\. (\d{1,2})\. (\d{1,2})\.", press).groups()
        press_date = f"{y}-{int(mo):02d}-{int(d):02d}"
    title = next((l for l in lines if re.search(r"\d{4}년 \d{1,2}월 수출입 운송비용 현황", l)), None)
    if not title:
        return insufficient("logistics", src, "자료 부족 — 보도자료 제목(자료 월)을 찾지 못함", "title not found")
    y, mo = map(int, re.search(r"(\d{4})년 (\d{1,2})월", title).groups())
    data_period = f"{y}-{mo:02d}"

    # 〔해상수출〕 표는 '구분 / 평균 운송비용 / 전월비 / 전년동월비' 순서의 문단으로 나열됨
    start = next((i for i, l in enumerate(lines) if l.startswith("□〔해상수출〕")), None)
    if start is None:
        return insufficient("logistics", src, "자료 부족 — 보도자료에서 '해상수출' 표를 찾지 못함", "section header not found")
    labels = ["미국서부", "미국동부", "유럽연합", "중동", "중국", "일본", "베트남"]
    # 다음 섹션(□〔해상수입〕 등) 직전까지만 읽는다. 넘어가면 수입 표 값으로 덮어써진다.
    end = next((j for j in range(start + 1, len(lines)) if lines[j].startswith("□〔")), len(lines))
    table = {}
    i = start
    while i < end - 3:
        if lines[i] in labels and re.fullmatch(r"[\d,]+천원", lines[i + 1]):
            table[lines[i]] = {
                "avg_cost_thousand_krw_per_2teu": int(lines[i + 1].replace("천원", "").replace(",", "")),
                "mom_pct": float(lines[i + 2].replace("%", "")),
                "yoy_pct": float(lines[i + 3].replace("%", "")),
            }
            i += 4
        else:
            i += 1
    if "미국서부" not in table or "미국동부" not in table:
        return insufficient("logistics", src, "자료 부족 — 미국 노선 해상수출 운임 값을 표에서 읽지 못함", f"parsed labels: {list(table)}")

    west, east = table["미국서부"], table["미국동부"]
    unit_note = next((l for l in lines if "2TEU" in l and "운송비용" in l), "")
    headline = f"{west['avg_cost_thousand_krw_per_2teu']:,}천원/2TEU · 전월비 {west['mom_pct']:+.1f}%"
    note = (
        f"한국→{TARGET_NAME} 서부 해상수출 평균 운송비용 {data_period}(관세청)"
        f" · 동부 {east['avg_cost_thousand_krw_per_2teu']:,}천원, 전월비 {east['mom_pct']:+.1f}%"
    )
    return {
        "key": "logistics",
        "headline": headline,
        "note": note,
        "state": "ok",
        "source": src,
        "as_of": month_end(y, mo),
        "data_class": "실제자료",
        "detail": {
            "press_release_date": press_date,
            "data_period": data_period,
            "direction": "해상수출(한국→해외)",
            "unit": "천원/2TEU (40피트 컨테이너 1개당 평균 총 운송비용, 할증료·수수료 포함)",
            "unit_note_raw": unit_note,
            "routes": table,
            # 종합 탭 스파크라인용 월별 운임 (보도자료 월별 표, 최근 13개월)
            "monthly": {
                route: [{"month": f"{yy}-{mm:02d}", "value": v} for yy, mm, v, *_ in (freight_monthly_series(lines, route) or [])[-13:]]
                for route in ("미국서부", "미국동부")
            },
            "parse_method": "hwpx(Contents/section0.xml) 문단 텍스트 정규식 추출",
        },
    }


# ---------------------------------------------------------------- 안정성
F_METHOD = ROOT / "config" / "stability_method.json"
F_FX = RAW / "external" / "fred_EXKOUS_krw_per_usd_monthly.csv"
F_FX_META = RAW / "external" / "fetch_meta.json"


def read_fred_csv(path):
    """FRED CSV -> {date_str: float}. 결측('.')은 건너뛴다."""
    out = {}
    with open(path, encoding="utf-8", newline="") as fh:
        rd = csv.reader(fh)
        hdr = next(rd)
        if hdr[:1] != ["observation_date"]:
            raise RuntimeError(f"unexpected header {hdr} in {path.name}")
        for r in rd:
            if len(r) >= 2 and r[1] not in ("", "."):
                out[r[0]] = float(r[1])
    return out


def monthly_means(daily, min_days, exclude_current):
    """일별 {date: v} -> [(y, m, mean, n)] 오름차순. 거래일이 min_days 미만인 달과 진행 중인 달은 제외."""
    buckets = {}
    for d, v in daily.items():
        y, m = int(d[:4]), int(d[5:7])
        buckets.setdefault((y, m), []).append(v)
    today = date.today()
    out = []
    for (y, m) in sorted(buckets):
        if exclude_current and (y, m) >= (today.year, today.month):
            continue
        vs = buckets[(y, m)]
        if len(vs) >= min_days:
            out.append((y, m, sum(vs) / len(vs), len(vs)))
    return out


def mean_abs_mom_pct(values):
    """전월비 등락률 절대값 평균(%). values 는 13개(12개 등락률)여야 정확하다."""
    chg = [abs(values[i] / values[i - 1] - 1) * 100 for i in range(1, len(values))]
    return sum(chg) / len(chg)


def volatility_metrics(series, win, hist_years, min_hist, labels):
    """series: [(y, m, value, ...)] 오름차순. 최근 win 개월 변동폭 + 과거 롤링 분포 대비 등급."""
    need = win + 1
    if len(series) < need:
        return None
    vals = [t[2] for t in series]
    recent = series[-need:]
    val = mean_abs_mom_pct(vals[-need:])
    hi = max(series[-win:], key=lambda t: t[2])
    lo = min(series[-win:], key=lambda t: t[2])
    mean_r = sum(vals[-win:]) / win
    cov = (sum((v - mean_r) ** 2 for v in vals[-win:]) / win) ** 0.5 / mean_r * 100
    recent_start = len(vals) - need
    hist_start = max(0, recent_start - hist_years * 12)
    rolling = sorted(mean_abs_mom_pct(vals[i:i + need]) for i in range(hist_start, recent_start))
    n = len(rolling)
    if n >= min_hist:
        t1, t2 = rolling[n // 3], rolling[(2 * n) // 3]
        grade = labels["low"] if val <= t1 else labels["mid"] if val <= t2 else labels["high"]
        cuts = [round(t1, 3), round(t2, 3)]
        pct_rank = round(sum(1 for x in rolling if x <= val) / n * 100, 1)
    else:
        grade, cuts, pct_rank = labels["none"], None, None
    return {
        "period": f"{recent[1][0]}-{recent[1][1]:02d}~{recent[-1][0]}-{recent[-1][1]:02d}",
        "last": (recent[-1][0], recent[-1][1]),
        "mean_abs_mom_pct": round(val, 3),
        "coefficient_of_variation_pct": round(cov, 3),
        "high": {"value": round(hi[2], 4), "month": f"{hi[0]}-{hi[1]:02d}"},
        "low": {"value": round(lo[2], 4), "month": f"{lo[0]}-{lo[1]:02d}"},
        "monthly": [{"month": f"{y}-{m:02d}", "value": round(v, 4)} for y, m, v, *_ in recent],
        "grading": {"history_windows": n, "tercile_cuts_pct": cuts, "percentile_rank_pct": pct_rank, "grade": grade},
    }


def freight_monthly_series(lines, route_label):
    """hwpx 문단에서 '해상 수출 운송비용 현황' 월별 표의 특정 노선(예: 미국서부) 금액 행을 읽는다."""
    start = next((i for i, l in enumerate(lines) if "해상 수출 운송비용 현황" in l and "∼" in l), None)
    if start is None:
        return None
    country, part = route_label[:2], route_label[2:]  # '미국' + '서부'
    i = start
    end = min(len(lines), start + 400)
    while i < end - 1:
        if lines[i] == country and lines[i + 1] == part:
            break
        i += 1
    else:
        return None
    i += 2
    series = []
    while i < end:
        if re.fullmatch(r"\d{4}", lines[i]) and i + 1 < end and lines[i + 1] == "금액":
            year = int(lines[i])
            j = i + 2
            month = 1
            while j < end and re.fullmatch(r"[\d,]+", lines[j]):
                series.append((year, month, float(lines[j].replace(",", "")), 1))
                month += 1
                j += 1
            i = j
        elif lines[i] in ("미국", "유럽연합", "중동", "중국", "일본", "베트남") and series:
            break  # 다음 노선
        else:
            i += 1
    return series or None


def extract_stability():
    h_market, h_stab = md5(F_WSTS), md5(F_WSTS_STAB)
    wsts_note = {
        "md5_marketability": h_market, "md5_stability": h_stab, "identical": h_market == h_stab,
        "note": "test1_stablity 폴더의 WSTS 파일은 시장성 파일과 동일하여 사용하지 않음",
    }
    if not F_METHOD.exists():
        return insufficient("stability", "", "자료 부족 — 안정성 산식 설정(junhee/config/stability_method.json) 없음", "method file missing", detail={"wsts": wsts_note})
    method = json.loads(F_METHOD.read_text(encoding="utf-8"))
    ind = method["indicator"]
    win = int(method["window"]["months"])
    excl = bool(method["window"].get("exclude_current_month", True))
    min_days = int(ind.get("min_trading_days_per_month", 10))
    hist_years = int(method["grading"]["history_years"])
    min_hist = int(method["grading"].get("min_history_windows", 36))
    labels = method["grading"]["labels"]
    ext = RAW / "external"
    fetched = json.loads(F_FX_META.read_text(encoding="utf-8")).get("fetched_at") if F_FX_META.exists() else None

    # (1) 원/달러 일별
    f_krw = ext / f"fred_{ind['krw_per_usd']}_daily.csv"
    if not f_krw.exists():
        return insufficient(
            "stability", f_krw.name,
            "자료 부족 — 원/달러 일별 환율 시계열 없음. `python junhee/scripts/fetch_fx.py` 로 수집 필요",
            "fx file missing", detail={"wsts": wsts_note},
        )
    krw_usd = read_fred_csv(f_krw)

    # (2) 통화별 원/외화 교차환율 → 월평균 → 변동폭
    per_ccy = {}
    for ccy, rule in ind["cross_rates"].items():
        if rule["op"] == "identity":
            daily = dict(krw_usd)
            src = f_krw.name
        elif rule["op"] == "direct":  # 이미 원/통화 로 제공되는 시계열 (ECOS)
            f = ext / rule["file"]
            if not f.exists():
                per_ccy[ccy] = {"state": "insufficient", "reason": f"{f.name} 없음 (fetch_fx_ecos.py 실행 필요, ECOS_API_KEY 환경변수)"}
                continue
            daily = read_fred_csv(f)
            src = f.name
        else:
            f = ext / f"fred_{rule['series']}_daily.csv"
            if not f.exists():
                per_ccy[ccy] = {"state": "insufficient", "reason": f"{f.name} 없음 (fetch_fx.py 실행 필요)"}
                continue
            other = read_fred_csv(f)
            if rule["op"] == "divide":
                daily = {d: krw_usd[d] / other[d] for d in krw_usd if d in other and other[d]}
            else:
                daily = {d: krw_usd[d] * other[d] for d in krw_usd if d in other}
            src = f"{f_krw.name} + {f.name}"
        months = monthly_means(daily, min_days, excl)
        met = volatility_metrics(months, win, hist_years, min_hist, labels)
        if met is None:
            per_ccy[ccy] = {"state": "insufficient", "reason": f"월 수 부족({len(months)})"}
            continue
        met.update({"state": "ok", "source": src, "desc": rule["desc"], "months_available": len(months)})
        per_ccy[ccy] = met

    # (3) 보조 지표: 해상수출 운임 월별 (미국서부)
    secondary = {"name": method["secondary"]["name"], "source": F_FREIGHT.name, "routes": {}}
    try:
        lines = hwpx_text_lines(F_FREIGHT)
        for country, route in method["secondary"]["route_by_country"].items():
            fs = freight_monthly_series(lines, route)
            if not fs:
                secondary["routes"][route] = {"state": "insufficient", "reason": "월별 표를 읽지 못함"}
                continue
            met = volatility_metrics(fs, win, hist_years, min_hist, labels)
            if met is None:
                secondary["routes"][route] = {"state": "insufficient", "reason": f"월 수 부족({len(fs)})"}
                continue
            met.update({"state": "ok", "months_available": len(fs), "unit": "천원/2TEU"})
            secondary["routes"][route] = met
    except Exception as e:  # noqa: BLE001
        secondary["error"] = repr(e)

    # (4) 목적국별 결과 (화면은 현재 목적국 것을 고른다)
    per_country = {}
    for country, spec in ind["currency_by_country"].items():
        ccy = spec["ccy"]
        if spec.get("available", True) is False or ccy not in per_ccy:
            per_country[country] = {
                "ccy": ccy, "state": "insufficient", "headline": "자료 부족",
                "note": f"{country} 결제통화 {ccy} 환율 시계열 없음 — {spec.get('note', '출처 확보 필요')}",
                "source": "", "as_of": None, "checked_on": date.today().isoformat(),
            }
            continue
        m = per_ccy[ccy]
        if m["state"] != "ok":
            per_country[country] = {"ccy": ccy, "state": "insufficient", "headline": "자료 부족",
                                    "note": f"{ccy} 환율 계산 불가 — {m['reason']}", "source": "", "as_of": None,
                                    "checked_on": date.today().isoformat()}
            continue
        grade = m["grading"]["grade"]
        grade_txt = f"과거 {hist_years}년 대비 {grade}" if m["grading"]["tercile_cuts_pct"] else grade
        headline = f"환율 평균 월간 변동폭 ±{m['mean_abs_mom_pct']:.1f}% · {grade_txt}"
        route = method["secondary"]["route_by_country"].get(country)
        r = secondary["routes"].get(route) if route else None
        sec_txt = f" · 보조 운임({route}) 변동폭 ±{r['mean_abs_mom_pct']:.1f}%({r['grading']['grade']}, {r['months_available']}개월)" if r and r.get("state") == "ok" else ""
        note = (
            f"원/{ccy} 월평균(연준 H.10 일별→월평균) {m['period']} · 최고 {m['high']['value']:,.2f}({m['high']['month']})"
            f" / 최저 {m['low']['value']:,.2f}({m['low']['month']}){sec_txt} · 산식 v{method['method_version']} {method['status']}"
        )
        ly, lm = m["last"]
        per_country[country] = {
            "ccy": ccy, "state": "ok", "headline": headline, "note": note,
            "source": m["source"], "as_of": month_end(ly, lm), "extra_note": spec.get("note"),
        }

    primary = per_country.get(TARGET_NAME)
    if not primary or primary["state"] != "ok":
        return insufficient(
            "stability", f_krw.name,
            primary["note"] if primary else f"자료 부족 — {TARGET_NAME} 통화 설정 없음",
            "primary country insufficient",
            detail={"per_country": per_country, "per_currency": per_ccy, "secondary": secondary, "wsts": wsts_note},
        )
    return {
        "key": "stability",
        "headline": primary["headline"],
        "note": primary["note"],
        "state": "ok",
        "source": primary["source"],
        "as_of": primary["as_of"],
        "data_class": "실제자료",
        "method_status": method["status"],
        "method_version": method["method_version"],
        "detail": {
            "indicator": ind["name"],
            "fetched_at": fetched,
            "window_months": win,
            "history_years": hist_years,
            "grading_method": method["grading"]["method"],
            "primary_country": TARGET_NAME,
            "per_country": per_country,
            "per_currency": per_ccy,
            "secondary": secondary,
            "method_file": F_METHOD.name,
            "decided": {"by": method.get("decided_by"), "on": method.get("decided_on")},
            "open_items": method.get("open_items", []),
            "wsts": wsts_note,
        },
    }


# ---------------------------------------------------------------- main
def main():
    items = []
    for fn in (extract_regulation, extract_market, extract_price, extract_logistics, extract_stability):
        key = fn.__name__.replace("extract_", "")
        try:
            item = fn()
        except Exception as e:  # noqa: BLE001  실패 시 값을 만들지 않고 실패로 기록
            item = insufficient(key, "", f"자료 부족 — 추출 실패 ({key})", f"{type(e).__name__}: {e}")
        items.append(item)
        (OUT / f"{key}.json").write_text(json.dumps(item, ensure_ascii=False, indent=2), encoding="utf-8")

    summary = {
        "schema_version": "0.2",
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "context": {"hs": HS_LIST, "hs_names": HS_NAMES, "country": TARGET_NAME, "company": None},
        "items": items,
    }
    (OUT / "dashboard_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"{'key':<11}{'state':<13}headline / note / source")
    for it in items:
        print(f"{it['key']:<11}{it['state']:<13}{it['headline']}")
        print(f"{'':<24}{it['note']}")
        print(f"{'':<24}source={it['source']} as_of={it['as_of']}")
    print("\nwritten:", OUT / "dashboard_summary.json")


if __name__ == "__main__":
    sys.exit(main())
