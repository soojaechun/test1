# (junhee) 2026-09-27 sanghyeob/analysis_market_price.py 에서 옮김. 바꾼 것: 패키지 import 경로·원본 자료 경로(junhee/data/raw)만. 계산 로직은 그대로.
"""Evidence for the existing company's market and price tabs; never a score.

API responses are fetched once by ``collect_trade``. ``evaluate_price`` reuses
its KCS observations and reads the supplied public files without changing them.
An optional requests-compatible session makes the entire pipeline testable offline.
"""

from collections import defaultdict
from copy import deepcopy
import csv
from datetime import date, datetime, timezone
from decimal import Decimal
from functools import lru_cache
from hashlib import sha256
from io import BytesIO, StringIO
import json
from pathlib import Path
import re
import xml.etree.ElementTree as ET

import requests

from .market_data import COMTRADE_URL, KCS_URL, _decimal, _fetch, _growth, _number
from .market_wsts import parse_wsts


COUNTRIES = {
    'US': {'name': '미국', 'reporter': 842, 'tariff': 'C840'},
    'CN': {'name': '중국', 'reporter': 156, 'tariff': 'C156'},
    'JP': {'name': '일본', 'reporter': 392, 'tariff': None},
    'DE': {'name': '독일', 'reporter': 276, 'tariff': 'U918'},
    'VN': {'name': '베트남', 'reporter': 704, 'tariff': None},
}
VALID = {'OBSERVED', 'OBSERVED_ZERO'}
ROOT = Path(__file__).resolve().parents[3]  # (junhee) 프로젝트 루트
PRICE_FILE = '(통계)2026년08월수출입물가지수무역지수(잠정).xlsx'
WSTS_FILE = 'WSTS-Historical-Billings-Report-Jul_2026.xlsx'
VERSION = 'company-market-price-v4-exact-observations'


def _now():
    return datetime.now(timezone.utc).isoformat()


def _month(offset, anchor):
    number = anchor.year * 12 + anchor.month - 1 + offset
    return f'{number // 12:04d}-{number % 12 + 1:02d}'


def _months(start, end):
    first = date.fromisoformat(start + '-01')
    n = (int(end[:4]) - first.year) * 12 + int(end[5:]) - first.month
    return [_month(i, first) for i in range(n + 1)]


def _scope(inputs):
    try:
        as_of = date.fromisoformat(inputs['as_of'])
        if inputs['country'] not in COUNTRIES or inputs['hs_edition'] != 'HS2022':
            raise ValueError
        if not re.fullmatch(r'[0-9]{6}', inputs['hs6']):
            raise ValueError
    except (ValueError, TypeError, KeyError):
        raise ValueError('국가·HS2022 6자리·분석 기준일을 확인해 주세요.') from None
    return as_of


def _obs(period, value=None, status='MISSING', evidence=None, reason=None, **extra):
    item = {'period': str(period), 'value': _number(value) if isinstance(value, Decimal) else value,
            'status': status, 'evidence': evidence or []}
    if isinstance(value, Decimal):
        item['value_decimal'] = str(value)
    if reason:
        item['reason'] = reason
    item.update(extra)
    return item


def _metric(identifier, label, observation, unit, formula):
    return {'id': identifier, 'label': label, 'value': observation.get('value'), 'unit': unit,
            'period': observation.get('period'), 'formula': formula,
            'status': observation['status'], 'evidence': observation.get('evidence', []),
            **({'reason': observation['reason']} if observation.get('reason') else {})}


def _factor(key, label, metrics, sources, warnings, series=None):
    count = sum(m['value'] is not None for m in metrics)
    return {'key': key, 'label': label, 'state': 'observed' if metrics and count == len(metrics)
            else 'partial' if count else 'insufficient', 'score': None,
            'score_status': 'PENDING_METHODOLOGY',
            'note': '원지표와 계산 근거입니다. 승인된 점수 산식이 없어 점수는 산출하지 않습니다.',
            'metrics': metrics, 'warnings': list(dict.fromkeys(warnings)),
            'checks': [{'label': '점수 산식', 'status': 'PENDING_METHODOLOGY',
                        'detail': '원값·비교기간·결측을 표시하며 성공확률이나 임의 점수로 변환하지 않습니다.'}],
            'sources': sources, 'series': series or []}


def _unavailable_source(source_id, provider, reason):
    return {'id': source_id, 'provider': provider, 'retrieved_at': _now(),
            'sha256': None, 'status': 'FETCH_ERROR', 'reason': reason,
            'adapter_version': VERSION}


def _batch_status(observations):
    states = {p['status'] for p in observations.values()}
    return 'OBSERVED' if states <= VALID else next(iter(states)) if len(states) == 1 else 'PARTIAL'


def _trade_batch(content, source, inputs, periods, frequency):
    """Reject an entire batch on an alien dimension, duplicate, or truncated count."""
    refs = [{'source_id': source['id']}]
    def failed(status, reason):
        return {p: _obs(p, status=status, evidence=refs, reason=reason) for p in periods}
    if content is None:
        return failed('FETCH_ERROR', source['reason'])
    try:
        doc = json.loads(content, parse_float=Decimal)
    except (ValueError, UnicodeDecodeError):
        return failed('INVALID_DATA', 'Comtrade JSON 응답을 읽을 수 없습니다.')
    if not isinstance(doc, dict) or not isinstance(doc.get('data'), list):
        return failed('INVALID_DATA', 'Comtrade 관측 목록이 없습니다.')
    if doc.get('error') or doc.get('errorMessage'):
        return failed('FETCH_ERROR', 'Comtrade 업무 오류입니다.')
    rows = doc['data']
    if (len(rows) > len(periods)
            or (doc.get('count') is not None and str(doc['count']) != str(len(rows)))):
        return failed('INVALID_DATA', '관측 건수·잘림 여부를 확인할 수 없습니다.')
    normalized = {str(p).replace('-', ''): p for p in periods}
    output = {p: _obs(p, evidence=refs, reason='해당 기간의 관측이 없습니다. 0으로 채우지 않습니다.') for p in periods}
    seen = set()
    expected = {'reporterCode': str(COUNTRIES[inputs['country']]['reporter']),
                'partnerCode': '0', 'partner2Code': '0', 'flowCode': 'M', 'cmdCode': inputs['hs6'],
                'classificationCode': 'H6', 'customsCode': 'C00', 'motCode': '0', 'freqCode': frequency}
    for index, row in enumerate(rows, 1):
        if (not isinstance(row, dict) or str(row.get('period')) not in normalized
                or any(str(row.get(k, '')) != value for k, value in expected.items())):
            return failed('INVALID_SCOPE', '반환 국가·기간·HS2022·수입·대세계 집계 차원이 요청과 다릅니다.')
        period = normalized[str(row['period'])]
        if period in seen:
            return failed('INVALID_DATA', '같은 기간의 Comtrade 관측이 중복되었습니다.')
        seen.add(period)
        evidence = [{'source_id': source['id'], 'record': index, 'field': 'primaryValue',
                     'period': period, 'hs6': inputs['hs6'], 'hs_edition': 'HS2022',
                     'reporter_code': row['reporterCode'], 'partner_code': 0, 'flow': 'M',
                     'valuation_basis': 'UNKNOWN',
                     'flags': {k: row.get(k) for k in ('isReported', 'isAggregate',
                               'isOriginalClassification', 'isNetWgtEstimated', 'isQtyEstimated')}}]
        try:
            value = _decimal(row.get('primaryValue'))
            output[period] = _obs(period, value, 'OBSERVED_ZERO' if value == 0 else 'OBSERVED', evidence)
        except ValueError:
            output[period] = _obs(period, status='INVALID_DATA', evidence=evidence,
                                  reason='수입금액이 없거나 유효하지 않습니다.')
    return output


def _kcs_batch(content, source, inputs, periods):
    refs = [{'source_id': source['id']}]
    def failed(status, reason):
        return {p: _obs(p, status=status, evidence=refs, reason=reason) for p in periods}
    if content is None:
        return failed('FETCH_ERROR', source['reason'])
    try:
        if b'<!DOCTYPE' in content.upper() or b'<!ENTITY' in content.upper():
            raise ET.ParseError
        doc = ET.fromstring(content)
    except ET.ParseError:
        return failed('INVALID_DATA', '관세청 XML 형식을 확인할 수 없습니다.')
    if doc.findtext('.//resultCode') != '00':
        return failed('FETCH_ERROR', '관세청 업무 응답이 정상 코드가 아닙니다.')
    groups = defaultdict(dict)
    for index, item in enumerate(doc.findall('.//item'), 1):
        row = {c.tag: (c.text or '').strip() for c in item}
        raw_period = row.get('year', '')
        if raw_period in ('총계', '합계', 'Total', 'TOTAL'):
            continue
        period = raw_period.replace('.', '-')
        code = row.get('hsCd', '')
        if (period not in periods or row.get('statCd') != inputs['country']
                or not re.fullmatch(r'[0-9]{6}(?:[0-9]{4})?', code)
                or not code.startswith(inputs['hs6'])):
            return failed('INVALID_SCOPE', '관세청 국가·연월·HS 범위가 요청과 다릅니다.')
        if code in groups[period]:
            return failed('INVALID_DATA', '관세청 월·품목 관측이 중복됩니다.')
        try:
            amount = _decimal(row.get('expDlr'))
        except ValueError:
            return failed('INVALID_DATA', '관세청 수출금액이 유효하지 않습니다.')
        try:
            weight = _decimal(row.get('expWgt'))
        except ValueError:
            weight = None
        groups[period][code] = (amount, weight, index)
    output = {}
    for period in periods:
        rows = groups[period]
        if not rows:
            output[period] = _obs(period, evidence=refs, reason='수출 관측월이 없습니다. 0으로 간주하지 않습니다.')
            continue
        details = {code: values for code, values in rows.items() if len(code) == 10}
        selected = details or rows
        amount = sum((v[0] for v in selected.values()), Decimal(0))
        if details and inputs['hs6'] in rows and rows[inputs['hs6']][0] != amount:
            output[period] = _obs(period, status='INVALID_DATA', evidence=refs,
                                  reason='HS6 소계와 HSK10 수출금액 합계가 다릅니다.')
            continue
        usable = {code: v for code, v in selected.items() if v[1] is not None and v[1] > 0}
        zero_zero = sum(v[0] == 0 and v[1] == 0 for v in selected.values())
        missing_weight = sum(v[0] > 0 and (v[1] is None or v[1] == 0) for v in selected.values())
        zero_value_weight = sum(v[0] == 0 and v[1] is not None and v[1] > 0 for v in selected.values())
        usable_amount = sum((v[0] for v in usable.values()), Decimal(0))
        weight = sum((v[1] for v in usable.values()), Decimal(0))
        evidence = [{'source_id': source['id'], 'record': v[2], 'period': period,
                     'field': 'expDlr/expWgt', 'hs_raw': code,
                     'value_usd': _number(v[0]), 'net_weight_kg': _number(v[1]),
                     'valuation_basis': 'FOB'} for code, v in selected.items()]
        output[period] = _obs(period, amount, 'OBSERVED_ZERO' if amount == 0 else 'OBSERVED', evidence,
                              usable_value_usd=_number(usable_amount), net_weight_kg=_number(weight),
                              records=len(selected), usable_records=len(usable),
                              zero_zero_records=zero_zero, missing_weight_records=missing_weight,
                              zero_value_positive_weight_records=zero_value_weight)
    return output


def _sum_periods(observations, periods, label):
    selected = [observations.get(p, _obs(p)) for p in periods]
    evidence = [r for obs in selected for r in obs.get('evidence', [])]
    invalid = next((p for p in selected if p['status'] not in VALID), None)
    if not selected or invalid:
        return _obs(label, status=invalid['status'] if invalid else 'MISSING', evidence=evidence,
                    reason='필요한 비교기간의 모든 관측값이 확인되지 않았습니다.')
    value = sum((Decimal(str(p.get('value_decimal', p['value']))) for p in selected), Decimal(0))
    return _obs(label, value, 'OBSERVED_ZERO' if value == 0 else 'OBSERVED', evidence)


def _change(current, previous, label):
    evidence = current['evidence'] + previous['evidence']
    if current['status'] not in VALID or previous['status'] not in VALID:
        invalid = current if current['status'] not in VALID else previous
        return _obs(label, status=invalid['status'], evidence=evidence,
                    reason='같은 범위의 현재·전년 비교자료가 필요합니다.')
    if previous['value'] == 0:
        return _obs(label, status='UNDEFINED', evidence=evidence, reason='분모가 0이므로 증가율을 정의할 수 없습니다.')
    value = Decimal(str(current.get('value_decimal', current['value']))) / Decimal(str(previous.get('value_decimal', previous['value']))) - 1
    return _obs(label, value, 'OBSERVED_ZERO' if value == 0 else 'OBSERVED', evidence)


def _industry(root_path, as_of):
    path = Path(root_path) / 'junhee' / 'data' / 'raw' / 'test1_marketability' / WSTS_FILE
    source_id = 'wsts-industry'
    try:
        content = path.read_bytes()
        parsed = parse_wsts(content, path.name)
    except (OSError, ValueError):
        return [], [], ['WSTS 산업 참고자료를 읽지 못했습니다. 국가 시장값을 대신하지 않습니다.']
    source = {'id': source_id, 'provider': 'WSTS', 'name': path.name, 'sha256': parsed['sha256'],
              'retrieved_at': _now(), 'observed_period': parsed['latest_period'],
              'excerpt': {'sheet': 'Monthly Data', 'unit_raw': '1000 US$', 'parser_version': parsed['parser_version']}}
    values = {p['period']: p for p in parsed['series'] if p['region'] == 'Worldwide'
              and p['period'] <= _month(-1, as_of)}
    if not values:
        return [], [source], ['기준일 이전의 WSTS 월별 관측이 없습니다.']
    latest = values[max(values)]
    observation = _obs(latest['period'], latest['value_usd'], 'OBSERVED',
                       [{'source_id': source_id, 'sheet': 'Monthly Data', 'cell': latest['source_cell'],
                         'field': 'Worldwide', 'value_raw': latest['value_raw'], 'unit_raw': '1000 US$'}])
    return [_metric('industry_world_sales', '세계 반도체 산업 매출(국가·HS 시장과 별도)', observation,
                    'USD', 'WSTS Worldwide 월별 원값 × 1,000')], [source], [
        'WSTS는 전체 반도체 산업 배경이며 국가·HS6 시장 또는 점수에 합산하지 않습니다.']


def _observation_index(rows, frequency):
    """Validate reusable observations without trusting a cached status alone."""
    output = {}
    pattern = r'[0-9]{4}' if frequency == 'A' else r'[0-9]{4}-[0-9]{2}'
    for row in rows:
        if not isinstance(row, dict):
            continue
        period = str(row.get('period', ''))
        if not re.fullmatch(pattern, period):
            continue
        try:
            date.fromisoformat(period + ('-01-01' if frequency == 'A' else '-01'))
        except ValueError:
            continue
        evidence = row.get('evidence', [])
        if period in output:
            output[period] = _obs(period, status='INVALID_DATA',
                                  evidence=output[period].get('evidence', []) + evidence,
                                  reason='같은 기준기간의 관측이 중복되어 사용할 수 없습니다.')
            continue
        observation = deepcopy(row)
        observation.setdefault('evidence', [])
        observation.setdefault('status', 'MISSING')
        observation.setdefault('value', None)
        if int(period[:4]) < 2022:
            observation.update(value=None, status='INCOMPATIBLE_SCOPE',
                               reason='HS2022 시행 전 자료는 자동 연결하지 않습니다.')
        elif observation['status'] in VALID:
            try:
                value = _decimal(observation.get('value_decimal', observation['value']))
                if _number(value) != observation['value']:
                    raise ValueError('관측의 정확값과 표시값이 다릅니다.')
                observation.update(value=_number(value), status='OBSERVED_ZERO' if value == 0 else 'OBSERVED')
            except ValueError:
                observation.update(value=None, status='INVALID_DATA', reason='관측 금액의 유효 범위를 확인할 수 없습니다.')
        else:
            observation['value'] = None
        output[period] = observation
    return output


def _annual_observation(annual, year):
    return annual.get(str(year), _obs(str(year),
        status='INCOMPATIBLE_SCOPE' if year < 2022 else 'MISSING',
        reason='동일 HS2022 범위의 기준연도 관측을 확인할 수 없습니다.'))


def rebase_market(trade, inputs, annual_year, latest_month):
    """Return a deep copy evaluated at the caller's exact common periods.

    No HTTP, file access, fallback or missing-value substitution occurs here.
    ``None`` leaves the associated indicators uncalculated. Raw observation
    lists and sources remain intact; only the market factor and annual_year
    are rebased. WSTS remains a separate reference with its original period.
    """
    as_of = _scope(inputs)
    if annual_year is not None and (isinstance(annual_year, bool) or not isinstance(annual_year, int)
                                    or annual_year < 1 or annual_year >= as_of.year):
        raise ValueError('완료된 연간 비교 기준을 확인해 주세요.')
    if latest_month is not None:
        try:
            if not isinstance(latest_month, str) or not re.fullmatch(r'[0-9]{4}-[0-9]{2}', latest_month):
                raise ValueError
            latest_date = date.fromisoformat(latest_month + '-01')
            if latest_month > _month(-1, as_of):
                raise ValueError
        except (ValueError, TypeError):
            raise ValueError('완료된 월별 비교 기준을 확인해 주세요.') from None
    result = deepcopy(trade)
    factor = result.get('factor', {})
    old_scope = factor.get('scope', {})
    if any(old_scope.get(field, inputs[field]) != inputs[field] for field in ('country', 'hs6', 'hs_edition')):
        raise ValueError('재계산 자료의 국가·HS 범위가 분석 조건과 다릅니다.')
    annual = _observation_index(result.get('annual_imports', []), 'A')
    monthly = _observation_index(result.get('monthly_imports', []), 'M')
    kcs = _observation_index(result.get('kcs_exports', []), 'M')
    missing = _obs('', reason='공통 비교 기준이 선택되지 않아 산출하지 않습니다.')
    missing['period'] = None
    import_value, exports, cagr = deepcopy(missing), deepcopy(missing), deepcopy(missing)
    ytd_change, recent_change = deepcopy(missing), deepcopy(missing)
    period_warnings = []
    if annual_year is not None:
        import_value = _annual_observation(annual, annual_year)
        annual_growth = {y: {'year': y, **_annual_observation(annual, y)}
                         for y in range(annual_year - 3, annual_year + 1)}
        cagr = _growth(annual_growth, annual_year, annual_year - 3, 3)
        cagr['period'] = f'{annual_year - 3}~{annual_year}'
        exports = _sum_periods(kcs, [f'{annual_year}-{m:02d}' for m in range(1, 13)], str(annual_year))
        if annual_year != as_of.year - 1:
            period_warnings.append(f'연간 비교 기준은 {annual_year}년입니다. 최근 완료 연도 {as_of.year - 1}년과 구분합니다.')
    else:
        period_warnings.append('공통 연간 비교 기준이 선택되지 않았습니다. 연간 지표는 미산출입니다.')
    if latest_month is not None:
        ytd = _months(f'{latest_date.year}-01', latest_month)
        prev_ytd = [f'{int(p[:4])-1:04d}{p[4:]}' for p in ytd]
        recent = [_month(i, latest_date) for i in range(-2, 1)]
        prior = [_month(i - 12, latest_date) for i in range(-2, 1)]
        ytd_change = _change(_sum_periods(monthly, ytd, latest_month),
                             _sum_periods(monthly, prev_ytd, latest_month),
                             f'{ytd[0]}~{latest_month} / {prev_ytd[0]}~{prev_ytd[-1]}')
        recent_change = _change(_sum_periods(monthly, recent, latest_month),
                                _sum_periods(monthly, prior, latest_month),
                                f'{recent[0]}~{recent[-1]} / {prior[0]}~{prior[-1]}')
        if latest_month != _month(-1, as_of):
            period_warnings.append(f'월별 비교 기준은 {latest_month}입니다. 요청 종료월 {_month(-1, as_of)}과 차이가 있습니다.')
    else:
        period_warnings.append('공통 월별 비교 기준이 선택되지 않았습니다. 과거 관측으로 임의 후퇴하지 않습니다.')
    metrics = [
        _metric('import_value', '대세계 연간 수입시장 규모', import_value, 'USD', '동일 국가·HS2022·HS6·연간·수입·대세계 primaryValue'),
        _metric('korea_exports', '한국의 해당국 연간 수출액', exports, 'USD', '총계·중복 소계를 제외한 12개월 expDlr 합계'),
        _metric('import_ytd_yoy', '수입액 누계 전년동기 증가율', ytd_change, 'ratio', '1월~선택 기준월 수입액 합 / 전년 같은 기간 합 − 1'),
        _metric('import_cagr3', '수입액 3년 CAGR', cagr, 'ratio', '(수입액(Y)/수입액(Y−3))^(1/3) − 1'),
        _metric('import_3m_yoy', '최근 3개월 수입액 전년동기 증가율', recent_change, 'ratio', '선택 기준월까지 3개월 금액 합 / 전년 같은 3개월 금액 합 − 1'),
    ]
    market_ids = {m['id'] for m in metrics}
    background = [m for m in factor.get('metrics', []) if m.get('id') not in market_ids]
    old_period_warnings = set(factor.get('period_warnings', []))
    warnings = [w for w in factor.get('warnings', []) if w not in old_period_warnings] + period_warnings
    series = [monthly[p] for p in sorted(monthly) if latest_month is not None and p <= latest_month]
    updated = _factor('market', '시장성', metrics, result.get('sources', factor.get('sources', [])), warnings, series)
    for check in factor.get('checks', []):
        if check not in updated['checks']:
            updated['checks'].append(check)
    updated['metrics'].extend(background)
    updated['period_warnings'] = period_warnings
    updated['scope'] = {'hs6': inputs['hs6'], 'hs_edition': 'HS2022', 'country': inputs['country'],
                        'annual_year': annual_year, 'requested_annual_year': as_of.year - 1,
                        'latest_month': latest_month, 'as_of': inputs['as_of']}
    result.update(factor=updated, annual_year=annual_year)
    return result


def collect_trade(inputs, keys, progress=None, session=None):
    """Read five annual and 60 calendar months of imports, plus 24+ KCS months.

    KCS also covers the preceding complete calendar year so a not-yet-published
    annual import series can fall back to the latest mutually observed year.
    """
    from urllib.parse import unquote
    as_of = _scope(inputs)
    own_session = session is None
    session = session or requests.Session()
    year = as_of.year - 1
    monthly_periods = [_month(i, as_of) for i in range(-60, 0)]
    kcs_periods = _months(f'{year - 1}-01', _month(-1, as_of))
    monthly, annual, kcs, sources = {}, {}, {}, []
    safe_keys = {name: keys.get(name, '') for name in ('UN_COMTRADE_API_KEY', 'KCS_TRADE_API_KEY')}

    def fetch(url, params, provider, sid, key_name):
        if not safe_keys[key_name]:
            return None, _unavailable_source(sid, provider, f'{provider} 인증키가 설정되지 않았습니다.')
        return _fetch(session, url, params, provider, sid, safe_keys)

    def comtrade(periods, frequency):
        sid = f'comtrade-{inputs["country"]}-{inputs["hs6"]}-{frequency}-{periods[0]}-{periods[-1]}'
        if progress:
            progress(f'{inputs.get("country_name", inputs["country"])} 수입 {periods[0]}~{periods[-1]} 확인')
        params = {'period': ','.join(p.replace('-', '') for p in periods),
                  'reporterCode': str(COUNTRIES[inputs['country']]['reporter']), 'cmdCode': inputs['hs6'],
                  'flowCode': 'M', 'partnerCode': '0', 'partner2Code': '0', 'customsCode': 'C00',
                  'motCode': '0', 'maxRecords': 50 if frequency == 'M' else 10, 'format': 'JSON',
                  'subscription-key': safe_keys['UN_COMTRADE_API_KEY']}
        url = COMTRADE_URL.replace('/C/A/', f'/C/{frequency}/')
        content, source = fetch(url, params, 'UN Comtrade', sid, 'UN_COMTRADE_API_KEY')
        source['observed_period'] = f'{periods[0]}~{periods[-1]}'
        result = _trade_batch(content, source, inputs, periods, frequency)
        source['status'] = _batch_status(result)
        sources.append(source)
        return result

    try:
        for y in range(year - 4, year + 1):
            period = str(y)
            annual.update(comtrade([period], 'A') if y >= 2022 else {
                period: _obs(period, status='INCOMPATIBLE_SCOPE', reason='HS2022 시행 전 연도는 자동 연결하지 않습니다.')})
        for y in sorted({p[:4] for p in monthly_periods}):
            periods = [p for p in monthly_periods if p[:4] == y]
            if int(y) < 2022:
                monthly.update({p: _obs(p, status='INCOMPATIBLE_SCOPE', reason='HS2022 시행 전 자료입니다.') for p in periods})
            else:
                monthly.update(comtrade(periods, 'M'))
        for y in sorted({p[:4] for p in kcs_periods}):
            periods = [p for p in kcs_periods if p[:4] == y]
            if int(y) < 2022:
                kcs.update({p: _obs(p, status='INCOMPATIBLE_SCOPE', reason='HS2022 시행 전 자료입니다.') for p in periods})
                continue
            sid = f'kcs-{inputs["country"]}-{inputs["hs6"]}-{periods[0]}-{periods[-1]}'
            if progress:
                progress(f'한국 수출 {periods[0]}~{periods[-1]} 금액·중량 확인')
            params = {'strtYymm': periods[0].replace('-', ''), 'endYymm': periods[-1].replace('-', ''),
                      'hsSgn': inputs['hs6'], 'cntyCd': inputs['country'],
                      'serviceKey': unquote(safe_keys['KCS_TRADE_API_KEY'])}
            content, source = fetch(KCS_URL, params, '관세청', sid, 'KCS_TRADE_API_KEY')
            source['observed_period'] = f'{periods[0]}~{periods[-1]}'
            observations = _kcs_batch(content, source, inputs, periods)
            source['status'] = _batch_status(observations)
            kcs.update(observations)
            sources.append(source)
    finally:
        if own_session:
            session.close()
    common_years = [y for y in range(year - 1, year + 1)
                    if annual[str(y)]['status'] in VALID
                    and _sum_periods(kcs, [f'{y}-{m:02d}' for m in range(1, 13)], str(y))['status'] in VALID]
    latest_states = {annual[str(year)]['status']} | {kcs[f'{year}-{m:02d}']['status'] for m in range(1, 13)}
    if common_years and latest_states <= VALID | {'MISSING', 'INCOMPATIBLE_SCOPE'}:
        year = max(common_years)
    warnings = ['수입은 목적국의 대세계 통계, 한국 수출은 한국 신고 FOB 통계입니다. 둘을 나눠 점유율을 만들지 않습니다.',
                'Comtrade 수입 primaryValue의 CIF/FOB 평가기준은 미확인입니다.',
                '기준일 이전 기간을 현재 조회한 수정 통계입니다. 당시 공표 자료만 사용한 역사적 재현은 아닙니다.',
                '누락·실패는 0이 아닙니다. 국가별 원지표는 별도 승인된 상대 비교 산식의 입력이며 성공확률이 아닙니다.']
    # The default individual view allows at most three months of publication
    # delay. A comparison caller must still select its own shared periods.
    valid_months = [p for p in monthly_periods if p >= _month(-4, as_of) and monthly[p]['status'] in VALID]
    latest = max(valid_months) if valid_months else None
    if latest is None:
        warnings.append('최근 완료월 대비 3개월 이내의 월별 관측을 확인하지 못했습니다.')
    background, industry_sources, industry_warnings = _industry(ROOT, as_of)
    sources.extend(industry_sources)
    factor = _factor('market', '시장성', background, sources, warnings + industry_warnings)
    factor['scope'] = {'hs6': inputs['hs6'], 'hs_edition': 'HS2022', 'country': inputs['country'],
                       'as_of': inputs['as_of']}
    trade = {'factor': factor, 'monthly_imports': [monthly[p] for p in monthly_periods],
             'annual_imports': [{'year': int(y), **annual[y]} for y in sorted(annual)],
             'kcs_exports': [kcs[p] for p in kcs_periods], 'annual_year': year, 'sources': sources}
    return rebase_market(trade, inputs, year, latest)


@lru_cache(maxsize=32)
def _read_tariff(path, hs6, mtime_ns, size):
    content = Path(path).read_bytes()
    reader = csv.DictReader(StringIO(content.decode('utf-8-sig')))
    required = {'reporter_code', 'partner_code', 'hs_code', 'year_dt', 'best_avlbl', 'imports'}
    if not required <= set(reader.fieldnames or []):
        raise ValueError('관세 CSV 열 구성이 다릅니다.')
    rows = []
    for index, row in enumerate(reader, 1):
        if row['hs_code'] == hs6:
            date.fromisoformat(row['year_dt'])
            rows.append((index, row))
    return sha256(content).hexdigest(), rows


def _tariff(inputs, root_path):
    code = COUNTRIES[inputs['country']]['tariff']
    empty = _obs(inputs['as_of'], status='NO_COVERAGE', reason='보유 WTO 파일에 해당 목적국·HS6 참고치가 없습니다.')
    if code is None:
        return empty, []
    path = Path(root_path) / 'junhee' / 'data' / 'raw' / 'test1_prices' / f'{code}_C410.csv'
    sid = f'wto-tariff-{code}-{inputs["hs6"]}'
    try:
        stat = path.stat()
        digest, rows = _read_tariff(str(path), inputs['hs6'], stat.st_mtime_ns, stat.st_size)
        if any(r['reporter_code'] != code or r['partner_code'] != 'C410' for _, r in rows):
            raise ValueError
        dates = [r['year_dt'] for _, r in rows]
        if len(dates) != len(set(dates)):
            raise ValueError
    except (OSError, ValueError, UnicodeError):
        return _obs(inputs['as_of'], status='INVALID_DATA', reason='보유 관세 CSV의 구조·범위를 확인할 수 없습니다.'), []
    eligible = [(i, r) for i, r in rows if r['year_dt'] <= inputs['as_of']]
    source = {'id': sid, 'provider': 'WTO–IMF Tariff Tracker', 'name': path.name,
              'retrieved_at': _now(), 'sha256': digest, 'observed_period': None,
              'basis': 'HS6_PROXY', 'applied_tariff_verified': False, 'weight_reference_year': None}
    if not eligible:
        empty['status'] = 'OUT_OF_PERIOD' if rows else 'NO_COVERAGE'
        empty['evidence'] = [{'source_id': sid}]
        return empty, [source]
    index, row = max(eligible, key=lambda pair: pair[1]['year_dt'])
    source.update(observed_period=row['year_dt'], excerpt={'data_record': index, **row})
    evidence = [{'source_id': sid, 'record': index, 'field': 'best_avlbl', 'action_date': row['year_dt'],
                 'hs6': inputs['hs6'], 'hs_edition': 'HS2022', 'reporter_code': code, 'partner_code': 'C410',
                 'basis': 'HS6_PROXY', 'applied_tariff_verified': False}]
    if not row['best_avlbl'].strip():
        return _obs(row['year_dt'], status='TARIFF_MISSING', evidence=evidence,
                    reason='기준일 이전의 최신 관세 행이 공란입니다. 과거 값으로 대체하지 않습니다.'), [source]
    try:
        value = _decimal(row['best_avlbl']) / 100
    except ValueError:
        return _obs(row['year_dt'], status='INVALID_DATA', evidence=evidence,
                    reason='관세 참고 세율이 유효하지 않습니다.'), [source]
    return _obs(row['year_dt'], value, 'REFERENCE_ONLY', evidence), [source]


@lru_cache(maxsize=4)
def _read_bok(path, mtime_ns, size):
    from openpyxl import load_workbook
    content = Path(path).read_bytes()
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        sheet = workbook['1.수출(기본분류)']
        rows = list(sheet.iter_rows(values_only=True))
        if '원화기준' not in str(rows[4][0]) or '2020=100' not in str(rows[4][6]):
            raise ValueError
        row_number, values = next((n, r) for n, r in enumerate(rows, 1)
                                  if str(r[0]).replace(' ', '') == '컴퓨터,전자및광학기기')
        observations = []
        for column in (2, 3):
            match = re.fullmatch(r'(\d{4})\.\s*(\d{1,2})(p?)', str(rows[6][column]).strip())
            if not match:
                raise ValueError
            observations.append((f'{match[1]}-{int(match[2]):02d}', _number(_decimal(values[column])),
                                 f'{chr(65 + column)}{row_number}', bool(match[3])))
        return sha256(content).hexdigest(), observations
    finally:
        workbook.close()


def _bok(root_path, as_of):
    path = Path(root_path) / 'junhee' / 'data' / 'raw' / 'test1_prices' / PRICE_FILE
    try:
        stat = path.stat()
        digest, observations = _read_bok(str(path), stat.st_mtime_ns, stat.st_size)
    except (OSError, ValueError, KeyError, StopIteration, ImportError):
        return [], [], ['한국은행 전자·광학 상위분류 가격지수를 확인하지 못했습니다.']
    sid = 'bok-electronics-export-price'
    source = {'id': sid, 'provider': '한국은행', 'name': path.name, 'retrieved_at': _now(),
              'sha256': digest, 'observed_period': f'{observations[0][0]}~{observations[-1][0]}',
              'excerpt': {'sheet': '1.수출(기본분류)', 'series': observations, 'base_year': 2020, 'basis': 'KRW'}}
    eligible = [r for r in observations if r[0] <= _month(-1, as_of)]
    if not eligible:
        return [], [source], ['기준일 이전 완결월의 한국은행 상위분류 지수가 없습니다.']
    period, value, cell, provisional = max(eligible)
    obs = _obs(period, value, 'PROVISIONAL' if provisional else 'OBSERVED',
               [{'source_id': sid, 'sheet': '1.수출(기본분류)', 'cell': cell, 'field': 'index'}])
    metric = _metric('electronics_export_price_index', '컴퓨터·전자·광학기기 수출물가지수(산업 배경)',
                     obs, 'index_2020_100', '원화 기준 원표 지수; 개별 제품 가격으로 환산하지 않음')
    return [metric], [source], ['한국은행 지수는 반도체·DRAM·HBM 전용 가격이 아닌 상위 산업 지수입니다.']


def _unit_value(observations, periods):
    selected = [observations.get(p, _obs(p)) for p in periods]
    refs = [r for p in selected for r in p.get('evidence', [])]
    period = f'{periods[0]}~{periods[-1]}'
    invalid = next((p for p in selected if p['status'] not in VALID), None)
    if invalid:
        return _obs(period, status=invalid['status'], evidence=refs,
                    reason='비교 12개월의 수출금액 관측이 모두 확인되지 않았습니다. ' + invalid.get('reason', ''))
    amount = sum((Decimal(str(p['usable_value_usd'])) for p in selected), Decimal(0))
    total = sum((Decimal(str(p['value'])) for p in selected), Decimal(0))
    weight = sum((Decimal(str(p['net_weight_kg'])) for p in selected), Decimal(0))
    records = sum(p['records'] for p in selected)
    usable_records = sum(p['usable_records'] for p in selected)
    zero_zero_records = sum(p.get('zero_zero_records', 0) for p in selected)
    missing_weight_records = sum(p.get('missing_weight_records', 0) for p in selected)
    zero_value_weight_records = sum(p.get('zero_value_positive_weight_records', 0) for p in selected)
    summary = {'source_id': refs[0]['source_id'] if refs else None,
               'source_ids': sorted({r['source_id'] for r in refs}),
               'kind': 'calculation_inputs', 'numerator_usd': _number(amount),
               'denominator_kg': _number(weight), 'total_value_usd': _number(total),
               'value_coverage_ratio': _number(amount / total) if total else None,
               'usable_records': usable_records, 'records': records,
               'excluded_records': records - usable_records - zero_zero_records,
               'zero_zero_records': zero_zero_records,
               'missing_weight_records': missing_weight_records,
               'zero_value_positive_weight_records': zero_value_weight_records,
               'excluded_value_usd': _number(total - amount),
               'months': len(periods), 'valuation_basis': 'FOB'}
    refs = refs + [summary]
    if weight <= 0:
        return _obs(period, evidence=refs, reason='확인된 양의 순중량이 없어 USD/kg 단가를 계산하지 않습니다.')
    complete = records == usable_records + zero_zero_records
    status = ('OBSERVED_ZERO' if amount == 0 else 'OBSERVED') if complete else 'PARTIAL_SAMPLE'
    reasons = []
    if zero_zero_records:
        reasons.append(f'금액·중량이 모두 0인 {zero_zero_records}행은 원본에 보존하며 합계에 기여하지 않습니다.')
    if not complete:
        reasons.append('중량이 확인된 표본만 사용했습니다. 전체 수출의 단가가 아니며 제외 금액을 추정 보충하지 않습니다.')
    if zero_value_weight_records:
        reasons.append('금액 0·양의 중량 관측이 포함됩니다. 무상거래 등 성격 미확인으로 판매가격을 뜻하지 않습니다.')
    return _obs(period, amount / weight, status, refs, reason=' '.join(reasons) or None)


def _company_prices(company, inputs, as_of):
    from .analysis_company_prices import evaluate_product_cost
    sheets = company.get('sheets', {})
    products = {str(p.get('제품ID')): p for p in sheets.get('제품정보', [])
                if re.sub(r'[.\s-]', '', str(p.get('HS코드', '')))[:6] == inputs['hs6']
                and str(p.get('HS버전', '')).replace(' ', '').upper() == 'HS2022'}
    country = inputs['country']
    country_names = {country, COUNTRIES[country]['name']}
    source_id = company.get('source', {}).get('id', 'company-upload')
    groups = defaultdict(list)
    exclusions = []
    start = _month(-12, as_of)
    end = _month(-1, as_of)
    for row in sheets.get('수출실적', []):
        product = str(row.get('제품ID', ''))
        if product not in products or row.get('목적국') not in country_names:
            continue
        if str(row.get('취소반품') or '').strip().upper() not in ('', 'N', 'NO', 'FALSE', '0', '정상', '아니오'):
            continue
        try:
            observed = date.fromisoformat(str(row.get('거래일'))[:10])
            if not start <= observed.strftime('%Y-%m') <= end:
                continue
            amount, quantity = _decimal(row.get('금액')), _decimal(row.get('수량'))
            if quantity <= 0 or not row.get('통화') or not row.get('단위'):
                raise ValueError
        except ValueError:
            exclusions.append(f'수출실적 {row.get("_row", "?")}행: 날짜·금액·수량·통화·단위를 확인해 주세요.')
            continue
        groups[(product, str(row['통화']).upper(), str(row['단위']))].append((row, amount, quantity))
    costs = defaultdict(list)
    for cost in sheets.get('원가·비용', []):
        costs[str(cost.get('제품ID', ''))].append(cost)
    metrics = []
    for (product, currency, unit), rows in sorted(groups.items()):
        amount = sum((r[1] for r in rows), Decimal(0))
        quantity = sum((r[2] for r in rows), Decimal(0))
        refs = [{'source_id': source_id, 'sheet': '수출실적', 'row': r[0].get('_row'),
                 'field': '금액/수량', 'amount': _number(r[1]), 'quantity': _number(r[2])} for r in rows]
        observation = _obs(f'{start}~{end}', amount / quantity, 'OBSERVED', refs)
        model = str(products[product].get('모델명', product))
        metrics.append(_metric(f'company_unit_price:{product}:{currency}:{unit}', f'{model} 기업 실적 단가',
                               observation, f'{currency}/{unit}', '선택 제품·목적국·동일 통화·단위의 금액 합 / 수량 합'))
        margin = evaluate_product_cost(costs[product], rows, product, currency, unit,
                                       period=f'{start}~{end}', source_id=source_id)
        basis = {'SCENARIO': '가정', 'COMPANY_COST_BASIS': '기업 기재 기간 기준'}.get(margin['status'], '확인 필요')
        formula = '(동일 통화 매출 − 제공 단위원가 × 수량) / 매출'
        formula += '; 기간 내 원가 일정 가정' if margin['status'] == 'SCENARIO' else '; 원가 적용기간에 모든 거래일 포함 여부 확인'
        metrics.append(_metric(f'company_product_margin:{product}:{currency}:{unit}', f'{model} 제품원가 차감률({basis})',
                               margin, 'ratio', formula))
    if not groups:
        metrics.append(_metric('company_unit_price', '선택 제품·목적국 기업 실적 단가',
                               _obs(f'{start}~{end}', reason='선택 제품·목적국에 유효한 동일 단위·통화의 수출실적이 없습니다.'),
                               None, '금액 합 / 수량 합'))
    for row in sheets.get('수출예정거래', []):
        product = str(row.get('제품ID', ''))
        if product not in products or row.get('목적국') not in country_names:
            continue
        refs = [{'source_id': source_id, 'sheet': row.get('_sheet', '수출예정거래'), 'row': row.get('_row'),
                 'field': '단가(USD)', 'transaction_id': row.get('거래ID')}]
        if isinstance(row.get('_field_origins'), dict):
            origins = row['_field_origins']
            refs[0]['field_origins'] = {field: origins[field] for field in
                                      ('단가(USD)', '희망판매단가', '통화', '단위', '작성기준일', '목적국') if field in origins}
            price_origin = origins.get('희망판매단가', origins.get('단가(USD)', {}))
            if price_origin.get('cell'):
                refs[0]['cell'] = price_origin['cell']
                refs[0]['row'] = price_origin['row']
                refs[0]['sheet'] = price_origin['sheet']
                refs[0]['field'] = price_origin['label']
        period = str(row.get('납기일') or row.get('작성기준일') or inputs['as_of'])
        currency, unit = str(row.get('통화') or '').upper(), str(row.get('단위') or '')
        obs = _obs(period, evidence=refs, reason='견적 단가·USD 통화·수량단위가 필요합니다. 빈칸은 0이 아닙니다.')
        if currency == 'USD' and unit:
            try:
                value = _decimal(row.get('단가(USD)'))
                obs = _obs(period, value, 'COMPANY_QUOTE', refs,
                           reason='기업 입력 견적입니다. 통계 USD/kg과 직접 비교하지 않으며 실제 체결·수익을 보증하지 않습니다.')
            except ValueError:
                pass
        metrics.append(_metric(f'company_quote:{row.get("거래ID", row.get("_row"))}',
                               f'{products[product].get("모델명", product)} 예정거래 견적 단가', obs,
                               f'{currency}/{unit}' if currency and unit else None,
                               '선택 제품·목적국 거래의 제공 단가; 여러 모델·거래를 평균하지 않음'))
    return metrics, exclusions


def evaluate_price(company, inputs, trade, root_path, keys, session=None):
    """Local references and uploaded company inputs; never repeat API collection."""
    as_of = _scope(inputs)
    rows = {p['period']: p for p in trade.get('kcs_exports', [])}
    periods = [_month(i, as_of) for i in range(-12, 0)]
    previous = [_month(i, as_of) for i in range(-24, -12)]
    current, prior = _unit_value(rows, periods), _unit_value(rows, previous)
    trend = _change(current, prior, f'{periods[0]}~{periods[-1]} / {previous[0]}~{previous[-1]}')
    tariff, tariff_sources = _tariff(inputs, root_path)
    sources = [s for s in trade.get('sources', []) if s.get('provider') == '관세청'] + tariff_sources
    warnings = ['통계 USD/kg 단가는 제품구성과 규격의 영향을 받습니다. 기업의 개당 견적·마진·경쟁력을 뜻하지 않습니다.',
                'WTO 값은 한국산 HS6 추정 참고 세율입니다. 세부세번·원산지·특혜·추가조치가 검증된 실제 적용세율이 아닙니다.',
                '관세는 분석 기준일의 파일 내 참고치입니다. 예정거래의 미래 적용세율을 확정하지 않습니다.',
                '관세 파일의 imports는 가중치용 참고금액입니다. 조치일의 시장규모나 성장률로 사용하지 않습니다.',
                '판매자·바이어 부담 및 청구가격·원가 포함 여부가 미확인된 비용은 0으로 처리하지 않습니다.']
    if inputs['country'] == 'DE':
        warnings.append('독일 관세 참고치는 EU 경제권(U918)의 공통 관세 파일입니다. 독일 자체 무역통계와 구분합니다.')
    metrics = [
        _metric('korea_export_unit_value', '한국의 해당국 수출 통계 단가', current, 'USD/kg', '금액·양의 순중량이 함께 있는 표본의 expDlr 합 / expWgt 합'),
        _metric('korea_export_unit_value_yoy', '수출 통계 단가 전년 동기간 변화율', trend, 'ratio', '최근 완결 12개월 합계단가 / 전년 같은 12개월 합계단가 − 1'),
        _metric('tariff_reference', '한국산 HS6 관세 참고 세율', tariff, 'ratio', '분석 기준일 이전의 가장 최근 조치행 best_avlbl / 100; 최신 공란은 유지'),
        _metric('applied_tariff', '거래별 실제 적용 관세', _obs(inputs['as_of'], status='NOT_VERIFIED',
                reason='목적국 세부세번·원산지·특혜·추가조치 검증이 필요합니다.'), 'ratio', '적용 원문·거래조건 검증 전 미산출'),
    ]
    reference = trade.get('world_reference', {})
    reference_scope = reference.get('scope', {})
    same_scope = all(reference_scope.get(k) == v for k, v in {
        'hs6': inputs['hs6'], 'hs_edition': 'HS2022', 'reporter': 'KR',
        'partner_scope': 'ALL_DESTINATIONS', 'flow': 'EXPORT', 'valuation_basis': 'FOB',
        'currency': 'USD', 'weight_unit': 'kg', 'as_of': inputs['as_of']}.items())
    sources.extend(reference.get('sources', []))
    warnings.extend(reference.get('warnings', []))
    if same_scope:
        reference_rows = _observation_index(reference.get('monthly_exports', []), 'M')
        world = _unit_value(reference_rows, periods)
    else:
        world = _obs(current['period'], status='INVALID_SCOPE' if reference_scope else 'MISSING',
                     reason='한국 전체 목적지의 동일 HS2022·기간·수출 FOB·USD·순중량 자료가 필요합니다.')
    relative = _obs(current['period'], evidence=current['evidence'] + world['evidence'],
                    reason='해당국과 전체 목적지의 같은 기간 완전 표본 단가가 모두 확인되어야 합니다.')
    if current['status'] in VALID and world['status'] in VALID and world['value'] > 0:
        relative = _obs(current['period'], Decimal(str(current['value'])) / Decimal(str(world['value'])),
                        'REFERENCE_ONLY', relative['evidence'],
                        reason='1이면 전체 목적지와 같은 통계 단가입니다. 제품구성 차이가 있으므로 기업 가격경쟁력·마진·점수가 아닙니다.')
    metrics.extend([
        _metric('korea_world_export_unit_value', '한국 전체 목적지 수출 통계 기준단가', world, 'USD/kg',
                '관세청 품목별 전체 수출의 같은 12개월 유효 금액 합 / 순중량 합'),
        _metric('relative_export_unit_value', '전체 목적지 대비 해당국 통계 단가 수준', relative, '배',
                '해당국 USD/kg 통계 단가 / 한국 전체 목적지 USD/kg 통계 단가'),
    ])
    for identifier, label, observation in (
            ('country_unit_value_coverage', '해당국 통계 단가 금액 포괄률', current),
            ('world_unit_value_coverage', '전체 목적지 통계 단가 금액 포괄률', world)):
        summary = next((e for e in observation['evidence'] if e.get('kind') == 'calculation_inputs'), None)
        if summary:
            metrics.append(_metric(identifier, label, _obs(observation['period'],
                summary['value_coverage_ratio'], 'REFERENCE_ONLY', [summary],
                reason=f'비기여 0금액·0중량 행 {summary["zero_zero_records"]}개; 중량을 확인할 수 없어 제외한 금액 {summary["excluded_value_usd"]} USD.'),
                'ratio', '중량이 확인된 관측의 금액 합 / 전체 수출금액 합; 총금액 0이면 정의하지 않음'))
    if relative['value'] is None:
        warnings.append('한국 전체 목적지 기준단가 또는 해당국 표본이 미완결되어 상대 단가 수준을 보류합니다.')
    warnings.append('통계 단가 표본 기준·세부 점수 산식 승인 전이므로 원지표를 점수로 변환하지 않습니다.')
    fx = trade.get('fx', {})
    metrics.extend(m for m in fx.get('metrics', []) if m.get('id', '').startswith('fx_reference_rate:'))
    sources.extend(fx.get('sources', []))
    warnings.extend(fx.get('warnings', []))
    bok, bok_sources, bok_warnings = _bok(root_path, as_of)
    metrics.extend(bok)
    sources.extend(bok_sources)
    company_metrics, company_warnings = _company_prices(company, inputs, as_of)
    metrics.extend(company_metrics)
    company_source = company.get('source')
    if company_source:
        sources.append({'provider': '기업 업로드', 'retrieved_at': None, **company_source})
    factor = _factor('price', '가격', metrics, sources, warnings + bok_warnings + company_warnings)
    factor['checks'].extend(fx.get('checks', []))
    factor['scope'] = {'hs6': inputs['hs6'], 'hs_edition': 'HS2022', 'country': inputs['country'],
                       'as_of': inputs['as_of'], 'price_period': current['period'], 'tariff_basis': 'HS6_PROXY'}
    return factor
