# -*- coding: utf-8 -*-
"""더미 기업 엑셀 2개 + junhee/data/raw 공개자료로 시연용 점수를 계산한다.

산식은 junhee/rules/demo_scoring.md 와 동일하다. 가상 데이터 · 시연용 산식이며 확정 채점 기준이 아니다.
출력: junhee/data/processed/companies/<company_id>.json, index.json

실행: python junhee/scripts/score_companies.py
"""
import calendar
import csv
import hashlib
import json
import re
import statistics
import sys
import warnings
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from extract_summary import read_wsts_worldwide  # noqa: E402

warnings.filterwarnings("ignore")
ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw"
SAMPLES = ROOT / "data" / "samples"
OUT = ROOT / "data" / "processed" / "companies"
OUT.mkdir(parents=True, exist_ok=True)

WEIGHTS = {"market": 35, "price": 30, "logistics": 20, "stability": 15}  # workspace.js defaults 와 동일
SERIES_MONTHS = 6
TRAIL = 12
DATA_CLASS = "가상 데이터 · 시연용 산식"
ISO2 = {"미국": "US", "중국": "CN", "영국": "GB", "인도": "IN", "캐나다": "CA", "멕시코": "MX", "스위스": "CH",
        "콜롬비아": "CO", "에콰도르": "EC", "캄보디아": "KH", "짐바브웨": "ZW", "EU": "EU", "독일": "DE", "일본": "JP"}
TARIFF_FILE = {"미국": "C840", "중국": "C156", "영국": "C826", "EU": "U918", "독일": "U918", "캐나다": "C124",
               "멕시코": "C484", "인도": "C356", "스위스": "C756", "콜롬비아": "C170", "에콰도르": "C218",
               "캄보디아": "C116", "짐바브웨": "C716"}
COMPANIES = [
    ("hanbit", "한빛반도체", "Hanbit Semiconductor", "한빛반도체_수출데이터_v0.1.xlsx"),
    ("daesung", "대성일렉트로닉스", "Daesung Electronics", "대성일렉트로닉스_수출데이터_v0.1.xlsx"),
]
FACTORS = [
    ("regulation", "규제 관문", "Regulation"),
    ("market", "시장성", "Marketability"),
    ("price", "가격", "Price & Tariff"),
    ("logistics", "물류", "Logistics"),
    ("stability", "안정성", "Stability"),
]
F_KOTRA = RAW / "test1_regulations" / "대한무역투자진흥공사_국별 대세계 수입규제 현황_20260603.csv"
F_CSL = RAW / "test1_regulations" / "ITA_consolidated_screening_list.csv"
F_HSK = RAW / "test1_regulations" / "HSK연계표_20260901.xlsx"
F_WSTS = RAW / "test1_marketability" / "WSTS-Historical-Billings-Report-Jul_2026.xlsx"


# ---------------------------------------------------------------- 유틸
def clip(x):
    return max(0.0, min(100.0, x))


def r1(x):
    return round(x + 1e-9, 1)


def ym_add(ym, k):
    y, m = ym
    n = y * 12 + (m - 1) + k
    return (n // 12, n % 12 + 1)


def ym_str(ym):
    return f"{ym[0]}-{ym[1]:02d}"


def month_end(ym):
    return date(ym[0], ym[1], calendar.monthrange(ym[0], ym[1])[1])


def as_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None  # 문자열 날짜(형식 오류)는 무효


def num(v):
    return float(v) if isinstance(v, (int, float)) else None


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def read_sheet(wb, name, id_col):
    """헤더 행(첫 셀 == id_col)을 찾아 dict 목록으로 읽는다."""
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r and r[0] == id_col)
    hdr = [h for h in rows[hi]]
    out = []
    for r in rows[hi + 1:]:
        if r and any(v is not None for v in r):
            out.append({hdr[i]: r[i] for i in range(min(len(hdr), len(r))) if hdr[i]})
    return out


def normalize_name(s):
    return re.sub(r"[^A-Z0-9]", "", str(s or "").upper())


# ---------------------------------------------------------------- 공개자료 로더 (캐시)
_cache = {}


def kotra_rows():
    if "kotra" not in _cache:
        with open(F_KOTRA, encoding="utf-8-sig", newline="") as fh:
            rows = list(csv.DictReader(fh))
        hs_cols = [h for h in rows[0] if "HS_코드" in h]
        _cache["kotra"] = (rows, hs_cols)
    return _cache["kotra"]


def csl_names():
    if "csl" not in _cache:
        names = set()
        with open(F_CSL, encoding="utf-8-sig", newline="") as fh:
            for r in csv.DictReader(fh):
                names.add(normalize_name(r["name"]))
                for alt in (r.get("alt_names") or "").split(";"):
                    if alt.strip():
                        names.add(normalize_name(alt))
        names.discard("")
        _cache["csl"] = names
    return _cache["csl"]


def hsk_candidates():
    """HS6 -> 통제번호가 있는 HSK 코드 수"""
    if "hsk" not in _cache:
        wb = openpyxl.load_workbook(F_HSK, read_only=True, data_only=True)
        cnt = defaultdict(int)
        for r in wb.worksheets[0].iter_rows(min_row=2, values_only=True):
            if r and r[0] and r[3]:
                cnt[str(r[0])[:6]] += 1
        _cache["hsk"] = dict(cnt)
    return _cache["hsk"]


def wsts_yoy_at(ym):
    """M 이전(포함) 가장 최근 WSTS 월의 전년동월비(%). 3개월 넘게 오래되면 None."""
    if "wsts" not in _cache:
        _cache["wsts"] = read_wsts_worldwide(F_WSTS)
    s = _cache["wsts"]
    cands = [k for k in s if k <= ym]
    if not cands:
        return None, None
    k = max(cands)
    if (ym[0] * 12 + ym[1]) - (k[0] * 12 + k[1]) > 3:
        return None, k
    prev = s.get((k[0] - 1, k[1]))
    if not prev:
        return None, k
    return (s[k] / prev - 1) * 100, k


def tariff_rate(country, hs6, ym):
    """목적국 관세 CSV 에서 hs6, year_dt <= M 말 인 최신 조치의 best_avlbl(%). 없으면 None."""
    code = TARIFF_FILE.get(country)
    if not code:
        return None
    key = ("tariff", code)
    if key not in _cache:
        f = RAW / "test1_prices" / f"{code}_C410.csv"
        if not f.exists():
            _cache[key] = None
        else:
            d = defaultdict(list)
            with open(f, encoding="utf-8-sig", newline="") as fh:
                for r in csv.DictReader(fh):
                    if r.get("best_avlbl", "").strip() == "":
                        continue  # 관세율이 비어 있는 행은 쓰지 않는다
                    d[r["hs_code"]].append((r["year_dt"], float(r["best_avlbl"])))
            _cache[key] = d
    table = _cache[key]
    if not table or hs6 not in table:
        return None
    lim = month_end(ym).isoformat()
    rows = [x for x in table[hs6] if x[0] <= lim]
    if not rows:
        return None
    return max(rows)[1]


# ---------------------------------------------------------------- 회사 데이터
def load_company(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    products = {p["제품ID"]: p for p in read_sheet(wb, "제품정보", "제품ID")}
    customers = {c["거래처ID"]: c for c in read_sheet(wb, "거래처", "거래처ID")}
    costs = {c["제품ID"]: c for c in read_sheet(wb, "원가·비용", "제품ID")}
    raw_actuals = read_sheet(wb, "수출실적", "실적ID")
    logistics = read_sheet(wb, "물류", "물류ID")
    # 유효 행 필터 + 완전 중복 제거
    seen, actuals = set(), []
    for r in raw_actuals:
        key = tuple(str(r.get(k)) for k in r)
        if key in seen:
            continue
        seen.add(key)
        d = as_date(r.get("거래일"))
        q, a = num(r.get("수량")), num(r.get("금액"))
        if d is None or str(r.get("취소반품") or "").strip().upper() == "Y" or not q or q <= 0 or not a or a <= 0:
            continue
        p = products.get(r.get("제품ID"))
        actuals.append({
            "id": r["실적ID"], "date": d, "ym": (d.year, d.month), "product": r.get("제품ID"),
            "hs6": str(p["HS코드"])[:6] if p else None, "customer": r.get("거래처ID"),
            "country": r.get("목적국"), "qty": q, "amount": a, "unit_price": a / q,
        })
    return {"products": products, "customers": customers, "costs": costs, "actuals": actuals,
            "logistics": logistics, "rows_total": len(raw_actuals), "rows_valid": len(actuals)}


def in_window(rows, ym, months=TRAIL):
    lo = ym_add(ym, -(months - 1))
    return [r for r in rows if lo <= r["ym"] <= ym]


def monthly_totals(rows, ym, months=TRAIL):
    t = defaultdict(float)
    for r in in_window(rows, ym, months):
        t[r["ym"]] += r["amount"]
    return t


def insufficient(note, inputs=None):
    return {"score": None, "state": "insufficient", "note": note, "inputs": inputs or {}}


# ---------------------------------------------------------------- 요인별 산식
def score_regulation(c, ym):
    rows = in_window(c["actuals"], ym)
    if not rows:
        return insufficient("자료 부족 — 추적 창에 유효한 수출실적이 없음")
    kotra, hs_cols = kotra_rows()
    pairs = sorted({(r["country"], r["hs6"]) for r in rows if r["hs6"]})
    kotra_hits = []
    for country, hs6 in pairs:
        iso = ISO2.get(country)
        for k in kotra:
            if k["규제시행국"].strip() == iso and k["한국대상여부"].strip() == "Y" and any(k[h].strip().startswith(hs6) for h in hs_cols if k[h].strip()):
                kotra_hits.append({"country": country, "hs6": hs6, "item": k["품목명"][:40]})
                break
    p1 = 30 if kotra_hits else 0
    names = csl_names()
    csl_hits = [cu["법인명"] for cu in c["customers"].values() if normalize_name(cu.get("법인명")) in names]
    p2 = 40 if csl_hits else 0
    cand = hsk_candidates()
    total = sum(r["amount"] for r in rows)
    flagged = sum(r["amount"] for r in rows if r["hs6"] and cand.get(r["hs6"], 0) > 0)
    share = flagged / total if total else 0
    p3 = 20 * share
    score = clip(100 - p1 - p2 - p3)
    hs_list = sorted({r["hs6"] for r in rows if r["hs6"]})
    return {
        "score": r1(score), "state": "ok",
        "needs_review": (p1 + p2 + p3) > 0,
        "note": (
            f"수입규제 해당 없음 · CSL 일치 없음 · 통제번호 후보 품목 비중 {share * 100:.0f}% (검토 필요, 해당 확정 아님)"
            if not kotra_hits and not csl_hits else
            f"수입규제 {len(kotra_hits)}건 · CSL 일치 {len(csl_hits)}건 · 통제번호 후보 비중 {share * 100:.0f}%"
        ),
        "inputs": {
            "hs6": hs_list, "countries": sorted({p[0] for p in pairs}),
            "kotra_hits": kotra_hits, "csl_hits": csl_hits,
            "hsk_candidate_share": round(share, 4), "hsk_codes_per_hs6": {h: cand.get(h, 0) for h in hs_list},
            "penalties": {"P1_kotra": p1, "P2_csl": p2, "P3_hsk": round(p3, 2)},
            "sources": [F_KOTRA.name, F_CSL.name, F_HSK.name],
        },
    }


def score_market(c, ym):
    yoy, wsts_month = wsts_yoy_at(ym)
    if yoy is None:
        return insufficient(f"자료 부족 — {ym_str(ym)} 기준 3개월 이내 WSTS 출하액 없음", {"wsts_latest": ym_str(wsts_month) if wsts_month else None})
    recent = sum(monthly_totals(c["actuals"], ym, 3).values())
    prev = sum(monthly_totals(c["actuals"], ym_add(ym, -3), 3).values())
    if prev <= 0:
        return insufficient("자료 부족 — 직전 3개월 수출실적이 없어 증가율을 계산할 수 없음")
    growth = (recent / prev - 1) * 100
    f, g = clip(50 + yoy / 2), clip(50 + growth)
    score = 0.5 * f + 0.5 * g
    return {
        "score": r1(score), "state": "ok",
        "note": f"세계 출하 전년동월비 {yoy:+.1f}%({ym_str(wsts_month)}) · 최근 3개월 수출액 {growth:+.1f}%",
        "inputs": {"wsts_yoy_pct": round(yoy, 2), "wsts_month": ym_str(wsts_month), "company_growth_3m_pct": round(growth, 2),
                   "recent_3m_usd": round(recent, 2), "prev_3m_usd": round(prev, 2), "f": r1(f), "g": r1(g),
                   "sources": [F_WSTS.name, "수출실적 시트"]},
    }


def score_price(c, ym):
    rows = in_window(c["actuals"], ym)
    if not rows:
        return insufficient("자료 부족 — 추적 창에 유효한 수출실적이 없음")
    total = sum(r["amount"] for r in rows)
    by_ch = defaultdict(float)
    for r in rows:
        by_ch[(r["country"], r["hs6"])] += r["amount"]
    weighted, missing = 0.0, []
    for (country, hs6), amt in by_ch.items():
        rate = tariff_rate(country, hs6, ym) if hs6 else None
        if rate is None:
            missing.append(f"{country} HS {hs6}")
            continue
        weighted += rate * amt / total
    if missing:
        return insufficient("자료 부족 — 관세율 없음: " + ", ".join(missing) + " (C<국가>_C410.csv 필요)")
    t = clip(100 - 2 * weighted)
    m_num, m_den, excluded = 0.0, 0.0, set()
    for r in rows:
        cost = c["costs"].get(r["product"])
        if not cost:
            excluded.add(r["product"])
            continue
        unit_cost = sum(num(cost.get(k)) or 0 for k in ("제품원가(단위당)", "운송비(단위당)", "보험료(단위당)", "포장비(단위당)"))
        m_num += (r["unit_price"] - unit_cost) / r["unit_price"] * 100 * r["amount"]
        m_den += r["amount"]
    if m_den <= 0:
        return insufficient("자료 부족 — 원가·비용 시트에 해당 제품 원가가 없음")
    margin = m_num / m_den
    m = clip(2.5 * margin)
    score = 0.5 * t + 0.5 * m
    return {
        "score": r1(score), "state": "ok",
        "note": f"가중 관세율 {weighted:.1f}% · 마진율 {margin:.1f}%" + (f" · 원가 없는 제품 제외 {len(excluded)}" if excluded else ""),
        "inputs": {"weighted_tariff_pct": round(weighted, 3), "margin_pct": round(margin, 2), "t": r1(t), "m": r1(m),
                   "tariff_by_channel": {f"{k[0]}/{k[1]}": tariff_rate(k[0], k[1], ym) for k in by_ch},
                   "excluded_products": sorted(excluded),
                   "sources": [f"{TARIFF_FILE[k[0]]}_C410.csv" for k in by_ch if k[0] in TARIFF_FILE] + ["수출실적 시트", "원가·비용 시트"]},
    }


def score_logistics(c, ym):
    lim = month_end(ym)
    lo = month_end(ym_add(ym, -TRAIL))
    by_id = {r["id"]: r for r in c["actuals"]}
    ontime, n, freight, linked = 0, 0, 0.0, 0.0
    for l in c["logistics"]:
        pa, aa = as_date(l.get("예정도착일")), as_date(l.get("실제도착일"))
        if pa is None or aa is None or not (lo < pa <= lim):
            continue
        n += 1
        ontime += 1 if aa <= pa else 0
        fr = num(l.get("운임(USD)")) or 0
        freight += fr
        deal = by_id.get(l.get("거래ID"))
        if deal:
            linked += deal["amount"]
    if n == 0:
        return insufficient("자료 부족 — 실제도착일이 있는 물류 행이 없음 (물류 시트 실제도착일 필요)")
    ontime_pct = ontime / n * 100
    if linked <= 0:
        return insufficient("자료 부족 — 물류 행이 수출실적(실적ID)과 연결되지 않아 운송비 비중을 계산할 수 없음")
    fr_pct = freight / linked * 100
    cscore = clip(100 - 10 * fr_pct)
    score = 0.7 * ontime_pct + 0.3 * cscore
    return {
        "score": r1(score), "state": "ok",
        "note": f"납기 준수율 {ontime_pct:.0f}% ({ontime}/{n}) · 운송비 비중 {fr_pct:.1f}%",
        "inputs": {"ontime_pct": round(ontime_pct, 1), "shipments": n, "ontime": ontime, "freight_usd": round(freight, 2),
                   "linked_amount_usd": round(linked, 2), "freight_share_pct": round(fr_pct, 2), "c": r1(cscore),
                   "sources": ["물류 시트", "수출실적 시트"]},
    }


def score_stability(c, ym):
    rows = in_window(c["actuals"], ym)
    totals = monthly_totals(c["actuals"], ym)
    if len(totals) < 6:
        return insufficient(f"자료 부족 — 최근 12개월 중 거래가 있는 달이 {len(totals)}개 (6개 이상 필요)")
    total = sum(r["amount"] for r in rows)

    def hhi(key):
        s = defaultdict(float)
        for r in rows:
            s[r[key]] += r["amount"]
        return sum((v / total) ** 2 for v in s.values())

    h_cust, h_ctry = hhi("customer"), hhi("country")
    h = (h_cust + h_ctry) / 2
    conc = clip((1 - h) * 100)
    vals = list(totals.values())
    cv = statistics.pstdev(vals) / statistics.mean(vals) * 100
    cvs = clip(100 - 2 * cv)
    score = 0.5 * conc + 0.5 * cvs
    return {
        "score": r1(score), "state": "ok",
        "note": f"거래처·목적국 집중도 HHI {h:.2f} · 월별 수출액 변동계수 {cv:.0f}% (기업 내부 거래 안정성, 국가위험 아님)",
        "inputs": {"hhi_customer": round(h_cust, 4), "hhi_country": round(h_ctry, 4), "hhi_mean": round(h, 4),
                   "cv_pct": round(cv, 2), "months_with_sales": len(totals), "conc": r1(conc), "cv_score": r1(cvs),
                   "sources": ["수출실적 시트", "거래처 시트"]},
    }


SCORERS = {"regulation": score_regulation, "market": score_market, "price": score_price,
           "logistics": score_logistics, "stability": score_stability}


def overall_of(factors):
    if any(factors[k]["state"] != "ok" for k in WEIGHTS):
        return None
    return sum(WEIGHTS[k] * factors[k]["score"] for k in WEIGHTS) / sum(WEIGHTS.values())


# ---------------------------------------------------------------- 회사 평가
def evaluate(company_id, name_ko, name_en, file_name):
    path = SAMPLES / file_name
    c = load_company(path)
    if not c["actuals"]:
        raise RuntimeError(f"{file_name}: 유효한 수출실적 없음")
    as_of = max(r["ym"] for r in c["actuals"])
    months = [ym_add(as_of, -k) for k in range(SERIES_MONTHS, -1, -1)]  # as_of-6 … as_of (전월 대비용 1개 더)
    per_month = {ym: {k: SCORERS[k](c, ym) for k, *_ in FACTORS} for ym in months}
    factors = []
    for key, ko, en in FACTORS:
        cur = per_month[as_of][key]
        prev = per_month[ym_add(as_of, -1)][key]
        series = [per_month[ym][key]["score"] for ym in months[1:]]
        f = {"key": key, "label_ko": ko, "label_en": en, "weight": WEIGHTS.get(key, 0) if key != "regulation" else None,
             "score": cur["score"], "state": cur["state"],
             "delta": r1(cur["score"] - prev["score"]) if cur["state"] == "ok" and prev["state"] == "ok" else None,
             "series": series, "series_months": [ym_str(ym) for ym in months[1:]],
             "note": cur["note"], "inputs": cur["inputs"]}
        if key == "regulation":
            f["gate"] = True
            f["needs_review"] = bool(cur.get("needs_review"))
        factors.append(f)
    fmap = {f["key"]: per_month[as_of][f["key"]] for f in factors}
    ov = overall_of(fmap)
    ov_prev = overall_of({k: per_month[ym_add(as_of, -1)][k] for k in fmap})
    ov_series = [overall_of({k: per_month[ym][k] for k in fmap}) for ym in months[1:]]
    overall = {"score": r1(ov) if ov is not None else None, "state": "ok" if ov is not None else "insufficient",
               "delta_vs_prev_month": r1(ov - ov_prev) if ov is not None and ov_prev is not None else None,
               "series": [r1(v) if v is not None else None for v in ov_series],
               "weights": WEIGHTS, "regulation_gate": "excluded",
               "needs_regulation_review": bool(fmap["regulation"].get("needs_review"))}

    # 주요 HS·목적국 (추적 창 수출액 기준)
    rows = in_window(c["actuals"], as_of)
    by_hs, by_ctry = defaultdict(float), defaultdict(float)
    for r in rows:
        by_hs[r["hs6"]] += r["amount"]
        by_ctry[r["country"]] += r["amount"]
    main_hs = max(by_hs, key=by_hs.get)
    main_ctry = max(by_ctry, key=by_ctry.get)

    oks = [f for f in factors if f["state"] == "ok" and f["key"] != "regulation"]
    best = max(oks, key=lambda f: f["score"]) if oks else None
    worst = min(oks, key=lambda f: f["score"]) if oks else None
    highlights = ""
    if best and worst:
        highlights = f"{best['label_ko']}({best['score']:.0f}점)이 가장 강하고 {worst['label_ko']}({worst['score']:.0f}점)이 가장 약합니다. "
    if overall["delta_vs_prev_month"] is not None:
        d = overall["delta_vs_prev_month"]
        highlights += f"종합 점수는 전월 대비 {d:+.1f}점 {'상승' if d > 0 else '하락' if d < 0 else '변화 없음'}했습니다. "
    if overall["needs_regulation_review"]:
        highlights += "규제 관문은 통제번호 후보 품목이 있어 원문 대조 검토가 필요합니다(해당 확정 아님). "
    highlights += "가상 데이터에 시연용 산식을 적용한 결과이며 실제 판정이 아닙니다."

    return {
        "company_id": company_id, "company_name": name_ko, "company_name_en": name_en,
        "file_name": file_name, "file_sha256": sha256(path),
        "as_of": month_end(as_of).isoformat(), "as_of_month": ym_str(as_of),
        "data_class": DATA_CLASS, "scoring_rule": "junhee/rules/demo_scoring.md (demo_scoring v0.1)",
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "hs": main_hs, "country": main_ctry, "country_iso2": ISO2.get(main_ctry),
        "hs_share": {k: round(v / sum(by_hs.values()), 3) for k, v in by_hs.items()},
        "country_share": {k: round(v / sum(by_ctry.values()), 3) for k, v in by_ctry.items()},
        "rows": {"actuals_total": c["rows_total"], "actuals_valid": c["rows_valid"], "logistics": len(c["logistics"])},
        "overall": overall, "factors": factors, "highlights": highlights,
    }


def main():
    index = {"generated_at": datetime.now().isoformat(timespec="seconds"), "data_class": DATA_CLASS,
             "scoring_rule": "junhee/rules/demo_scoring.md", "companies": []}
    results = []
    for cid, ko, en, fn in COMPANIES:
        res = evaluate(cid, ko, en, fn)
        (OUT / f"{cid}.json").write_text(json.dumps(res, ensure_ascii=False, indent=2), encoding="utf-8")
        index["companies"].append({k: res[k] for k in ("company_id", "company_name", "company_name_en", "file_name", "file_sha256", "as_of", "hs", "country", "country_iso2")} | {"overall_score": res["overall"]["score"]})
        results.append(res)
    (OUT / "index.json").write_text(json.dumps(index, ensure_ascii=False, indent=2), encoding="utf-8")

    # 콘솔 표
    print(f"{'항목':<8}" + "".join(f"{r['company_name']:>22}" for r in results))
    for i, (key, ko, _) in enumerate(FACTORS):
        cells = []
        for r in results:
            f = r["factors"][i]
            cells.append(f"{f['score']:.1f} ({f['delta']:+.1f})" if f["state"] == "ok" else "자료 부족")
        print(f"{ko:<8}" + "".join(f"{c:>22}" for c in cells))
    print(f"{'종합':<8}" + "".join(f"{(str(r['overall']['score']) + ' (' + format(r['overall']['delta_vs_prev_month'], '+.1f') + ')') if r['overall']['score'] is not None else '자료 부족':>22}" for r in results))
    print(f"{'기준월':<8}" + "".join(f"{r['as_of_month']:>22}" for r in results))
    print(f"{'규제검토':<8}" + "".join(f"{str(r['overall']['needs_regulation_review']):>22}" for r in results))
    print("written:", OUT)


if __name__ == "__main__":
    main()
